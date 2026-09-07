import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_row(row, bg_color=None, is_header=False, font_size=9.5):
    for cell in row.cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, top=90, bottom=90, left=130, right=130)
        if bg_color:
            set_cell_background(cell, bg_color)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            for r in p.runs:
                r.font.name = "Segoe UI"
                r.font.size = Pt(font_size)
                if is_header:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    r.font.color.rgb = RGBColor(15, 23, 42)

def add_heading_1(doc, text):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    for r in h.runs:
        r.font.name = "Segoe UI"
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
    return h

def add_heading_2(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    for r in h.runs:
        r.font.name = "Segoe UI"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)
    return h

def add_heading_3(doc, text):
    h = doc.add_heading(text, level=3)
    h.paragraph_format.space_before = Pt(11)
    h.paragraph_format.space_after = Pt(3)
    h.paragraph_format.keep_with_next = True
    for r in h.runs:
        r.font.name = "Segoe UI"
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(37, 99, 235)
    return h

def create_overview_doc():
    doc = docx.Document()

    # A4 Margins (1 inch)
    sec = doc.sections[0]
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_t = title_p.add_run("Aequs EduTrack — Project Timeline & Visual Overview")
    run_t.font.name = "Segoe UI"
    run_t.font.size = Pt(22)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    run_s = sub_p.add_run("Executive Engineering Roadmap, Visual Gantt Charts & Architecture Milestones")
    run_s.font.name = "Segoe UI"
    run_s.font.size = Pt(11)
    run_s.font.color.rgb = RGBColor(71, 85, 105)

    # 1. Executive Snapshot
    add_heading_2(doc, "Executive Snapshot")
    snapshot_items = [
        ("Project Name", "Aequs EduTrack (Educational CSR Management Platform)"),
        ("Active Timeframe", "August 12, 2026 – September 07, 2026"),
        ("Project Duration", "27 Calendar Days (14 Active Development Days)"),
        ("Total Git Commits", "61 Commits (+ 1 Architectural Optimization Sprint)"),
        ("Core Engineering Team", "2 Primary Contributors (Rishab / Code-Cool-2006, Saloni Dalvi)"),
        ("Rishab / Code-Cool-2006", "42 Commits (69%) — Lead Architecture, UI/UX, School/Volunteer Portals, Testing, DevOps"),
        ("Saloni Dalvi / salonidalvi-008", "19 Commits (31%) — Core Scaffolding, Distributions, CourseMaster, Internship Portal, Eligibility"),
        ("Automated Test Suite", "78 / 78 Automated Tests Passing (100% OK in 5.6s)"),
        ("Django System Health", "0 Issues, 0 Silenced Warnings (Clean Architecture)"),
    ]

    t_snap = doc.add_table(rows=len(snapshot_items), cols=2)
    t_snap.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_snap, color="E2E8F0", sz="6")
    for idx, (k, v) in enumerate(snapshot_items):
        r = t_snap.rows[idx]
        r.cells[0].width = Inches(2.2)
        r.cells[1].width = Inches(4.07)
        r.cells[0].paragraphs[0].add_run(k).bold = True
        r.cells[1].paragraphs[0].add_run(v)
        bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        format_row(r, bg_color=bg, is_header=False, font_size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 2. Visual Roadmap & Gantt Chart
    add_heading_2(doc, "Visual Roadmap & Gantt Chart")
    p_g = doc.add_paragraph("The visual Gantt chart illustrates the end-to-end multi-week timeline across all 10 engineering phases from initial project scaffolding to architectural audit:")
    p_g.paragraph_format.space_after = Pt(8)

    gantt_path = "docs/images/gantt_chart.png"
    if os.path.exists(gantt_path):
        doc.add_picture(gantt_path, width=Inches(6.27))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.space_after = Pt(14)

    # 3. Module Evolution & Feature Maturation
    add_heading_2(doc, "Module Evolution & Feature Maturation")
    p_m = doc.add_paragraph("The lifecycle diagram below traces the release and maturation milestones across the core functional modules of Aequs EduTrack:")
    p_m.paragraph_format.space_after = Pt(8)

    module_path = "docs/images/module_evolution.png"
    if os.path.exists(module_path):
        doc.add_picture(module_path, width=Inches(6.27))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.space_after = Pt(14)

    # 4. Development Phases & Commit Distribution
    add_heading_2(doc, "Development Phases & Commit Distribution")
    p_p = doc.add_paragraph("Distribution of development activity across phases alongside total team commit share:")
    p_p.paragraph_format.space_after = Pt(8)

    phases_path = "docs/images/phases_summary.png"
    if os.path.exists(phases_path):
        doc.add_picture(phases_path, width=Inches(6.27))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.space_after = Pt(14)

    # 5. Phase-by-Phase Chronological Log
    add_heading_2(doc, "Phase-by-Phase Chronological Log")
    doc.add_paragraph("Detailed breakdown of engineering deliverables, architectural milestones, and commit activity across all 10 phases:").paragraph_format.space_after = Pt(6)

    phases_log = [
        {
            "phase": "Phase 1: Project Scaffolding & Core Architecture",
            "dates": "August 12, 2026",
            "commits": "2 commits",
            "lead": "Saloni Dalvi",
            "items": [
                "Initialized Django project with custom User model and 5 RBAC roles (Super Admin, Admin, Accountant, Volunteer, Event Coordinator).",
                "Core database models established: School (UDISE codes, headmasters), Student (demographics), AcademicRecord (marks, promotion), EligibilityRecord (5 benefit types), and Laptop / LaptopAssignment tracking.",
                "Basic list and creation forms for students, schools, and benefit distributions.",
            ]
        },
        {
            "phase": "Phase 2: REST APIs & Management Services",
            "dates": "August 13, 2026",
            "commits": "1 commit",
            "lead": "Rishab",
            "items": [
                "Warehouse models: InventoryItem and StockTransaction for stock-in/stock-out management.",
                "Interactive Academics Portal and Inventory Portal templates.",
                "Automated CLI seed commands: seed_schools, seed_inventory, seed_academics.",
                "Multi-item StudyKit services layer.",
            ]
        },
        {
            "phase": "Phase 3: Major Feature Sprint, UI/UX Overhaul & Production Deployment",
            "dates": "August 18, 2026",
            "commits": "19 commits",
            "lead": "Rishab, Saloni Dalvi",
            "items": [
                "Added Internship module (programs, candidate placements, milestones).",
                "Foundational scaffolding for Events, Volunteers, and Reports.",
                "Sidebar Evolution: 3 iterations (dark -> unified dark -> high-contrast royal blue/white light theme).",
                "Global Design System: Unified typography, custom CSS tokens, compact data density, and shadcn-style card aesthetics.",
                "Real-world Belagavi district mock data seeder.",
                "Production Cloud Deployment: Docker containerization, Render PaaS blueprint, Gunicorn WSGI, Whitenoise static streaming, and PostgreSQL database support.",
            ]
        },
        {
            "phase": "Phase 4: Authentication & Bulk Data Operations",
            "dates": "August 24, 2026",
            "commits": "3 commits",
            "lead": "Rishab, Saloni Dalvi",
            "items": [
                "Branded login/logout portal featuring Aequs SEZ background and glassmorphism styling.",
                "Session protection and role authentication guards across all portal views.",
                "CSV Bulk Upload engine supporting Students, Schools, Academics, and Inventory.",
            ]
        },
        {
            "phase": "Phase 5: Academic Enhancements & Enterprise SQL Server",
            "dates": "August 25–27, 2026",
            "commits": "6 commits",
            "lead": "Saloni Dalvi, Rishab",
            "items": [
                "CourseMaster model for dynamic curriculum streams and grades.",
                "Configured Microsoft SQL Server (mssql-django / pyodbc) for on-prem enterprise deployments.",
                "Completed full Internship Management Portal with milestone tracking and stipend evaluation.",
            ]
        },
        {
            "phase": "Phase 6: School Management Portal & Automated Test Suite",
            "dates": "August 28, 2026",
            "commits": "6 commits",
            "lead": "Rishab, Saloni Dalvi",
            "items": [
                "Dedicated School Portal dashboard with grade-by-grade student strength matrix (Classes 1–10).",
                "Engineered comprehensive 78-case automated test suite validating models, permissions, and REST endpoints.",
            ]
        },
        {
            "phase": "Phase 7: Eligibility Engine Fixes & Class Resolvers",
            "dates": "August 31, 2026",
            "commits": "5 commits",
            "lead": "Saloni Dalvi, Rishab",
            "items": [
                "Fixed critical eligibility bug where higher-education degrees (BE, Diploma, ITI, PUC) were conflated with K-10 school grades.",
                "Added resilient class name resolver (10th <-> Class 10, PUC <-> 12th).",
                "Published initial User Manual and Timeline documentation.",
            ]
        },
        {
            "phase": "Phase 8: Inventory Logistics, Resource Requests & PR #1",
            "dates": "September 01–02, 2026",
            "commits": "5 commits",
            "lead": "Rishab, Saloni Dalvi",
            "items": [
                "Event Resource Requests: Event coordinators can request warehouse equipment directly.",
                "Repaired inventory bulk upload with atomic transaction handling.",
                "Stock status indicators (In Stock, Low Stock, Depleted) in inventory views.",
                "Merged Pull Request #1 into main codebase.",
            ]
        },
        {
            "phase": "Phase 9: Volunteer Management System, QR Generation & Public Registration",
            "dates": "September 04–05, 2026",
            "commits": "14 commits",
            "lead": "Rishab, Saloni Dalvi",
            "items": [
                "Complete Volunteer Portal: Volunteer directory, hours contributed, event assignments, and activity logs.",
                "Tokenized VolunteerFormLink generator allowing administrators to share secure external signup links.",
                "Live dynamic PNG QR code generation (Reed-Solomon Level H error correction).",
                "Public, standalone mobile-responsive volunteer registration page with activity audit logging.",
                "Direct CSV and Excel export endpoints for volunteer rosters.",
            ]
        },
        {
            "phase": "Phase 10: Architectural Audit & Codebase Streamlining",
            "dates": "September 07, 2026",
            "commits": "1 sprint",
            "lead": "Rishab",
            "items": [
                "Purged 17 dead backup and text dump files (project_frontend.txt, *_backup.py, *.corrupt-backup).",
                "Removed 5 empty skeleton apps (audit, documents, notifications, donations, expense_management) from disk and INSTALLED_APPS.",
                "Unified duplicate QR generation views into a single reusable handler.",
                "Removed redundant transitive dependencies (asgiref, sqlparse) from requirements.txt.",
                "Net Impact: -29,811 lines of dead code removed, 78/78 tests passing (100% green), 0 system check issues.",
            ]
        },
    ]

    for p_info in phases_log:
        add_heading_3(doc, p_info["phase"])
        meta_p = doc.add_paragraph()
        r1 = meta_p.add_run("Dates: ")
        r1.bold = True
        meta_p.add_run(f"{p_info['dates']} | ")
        r2 = meta_p.add_run("Volume: ")
        r2.bold = True
        meta_p.add_run(f"{p_info['commits']} | ")
        r3 = meta_p.add_run("Lead: ")
        r3.bold = True
        meta_p.add_run(p_info['lead'])
        meta_p.paragraph_format.space_after = Pt(2)

        for itm in p_info["items"]:
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.space_before = Pt(1)
            bp.paragraph_format.space_after = Pt(2)
            rb = bp.add_run(itm)
            rb.font.name = "Segoe UI"
            rb.font.size = Pt(9)
            rb.font.color.rgb = RGBColor(30, 41, 59)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 6. Contributor Breakdown
    add_heading_2(doc, "Contributor Breakdown (61 Total Commits)")
    doc.add_paragraph("The repository history reflects contributions from two primary software engineers:").paragraph_format.space_after = Pt(4)

    contrib_summary = [
        ("Rishab / Code-Cool-2006", "42", "69%", "Lead Architect, UI/UX & Systems Engineer",
         "• REST APIs & management commands (seed_schools, seed_inventory, seed_academics)\n"
         "• UI/UX design system (3 sidebar iterations, global CSS, shadcn card aesthetics)\n"
         "• School Management Portal, Academics Portal, and Inventory Portal\n"
         "• Automated 78-case unit/API test suite (100% pass rate in 5.6s)\n"
         "• Volunteer Management System (CRUD, public signup forms, live QR generation)\n"
         "• Production deployment blueprints (Docker, Render PaaS, Gunicorn, Whitenoise)\n"
         "• Codebase optimization (-29,811 dead lines purged across 52 files)"),
        ("Saloni Dalvi / salonidalvi-008", "19", "31%", "Core Services & Logistics Engineer",
         "• Initial project foundation and core schema scaffolding (13 apps)\n"
         "• Benefit distribution services and Study Kit assembly workflows\n"
         "• Dynamic CourseMaster and academic grade/stream filters\n"
         "• Enterprise Microsoft SQL Server (mssql-django / pyodbc) integration\n"
         "• Full Internship Management Portal (company tracks, milestones, stipends)\n"
         "• Eligibility bugfix (course separation & class equivalent resolution)\n"
         "• Bulk CSV onboarding pipelines across models"),
    ]

    t_cb = doc.add_table(rows=len(contrib_summary) + 1, cols=4)
    t_cb.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_cb, color="E2E8F0", sz="4")

    h_cb = t_cb.rows[0]
    h_cb.cells[0].width = Inches(1.8)
    h_cb.cells[1].width = Inches(0.7)
    h_cb.cells[2].width = Inches(0.7)
    h_cb.cells[3].width = Inches(3.07)
    h_cb.cells[0].paragraphs[0].add_run("Contributor")
    h_cb.cells[1].paragraphs[0].add_run("Commits")
    h_cb.cells[2].paragraphs[0].add_run("Share")
    h_cb.cells[3].paragraphs[0].add_run("Functional Domain & Impact")
    format_row(h_cb, bg_color="1E3A8A", is_header=True, font_size=8.5)

    for idx, (c_name, c_cnt, c_pct, c_role, c_imp) in enumerate(contrib_summary):
        r = t_cb.rows[idx + 1]
        r.cells[0].width = Inches(1.8)
        r.cells[1].width = Inches(0.7)
        r.cells[2].width = Inches(0.7)
        r.cells[3].width = Inches(3.07)
        r.cells[0].paragraphs[0].add_run(c_name).bold = True
        r.cells[1].paragraphs[0].add_run(c_cnt)
        r.cells[2].paragraphs[0].add_run(c_pct).bold = True
        r.cells[3].paragraphs[0].add_run(f"{c_role}\n{c_imp}")
        bg = "F8FAFC" if idx % 2 == 0 else "FFFFFF"
        format_row(r, bg_color=bg, is_header=False, font_size=8)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Footer
    foot_p = doc.add_paragraph()
    foot_p.paragraph_format.space_before = Pt(16)
    foot_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf = foot_p.add_run("Aequs EduTrack Project Timeline & Visual Overview • September 07, 2026\n© 2026 Aequs Foundation. All rights reserved.")
    rf.font.name = "Segoe UI"
    rf.font.size = Pt(8.5)
    rf.font.color.rgb = RGBColor(100, 116, 139)

    out_file = "Aequs EduTrack — Project Timeline & Visual Overview.docx"
    doc.save(out_file)
    print(f"Successfully generated: {out_file}")

if __name__ == "__main__":
    create_overview_doc()
