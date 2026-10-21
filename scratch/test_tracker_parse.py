import os
import datetime
import pandas as pd
import openpyxl

def test_parse(file_path, academic_year):
    print(f"=== Testing {file_path} ({academic_year}) ===")
    wb = openpyxl.load_workbook(file_path, read_only=True)
    start_year = int(academic_year.split('-')[0])
    
    # 1. Weekly Tracker sheets
    tracker_sheets = [s for s in wb.sheetnames if 'tracker' in s.lower() or s.strip() in ['Others', 'Winter & Summer Camp', 'SAP']]
    print("Found tracker sheets:", tracker_sheets)
    total_sessions = 0
    for s in tracker_sheets:
        df = pd.read_excel(file_path, sheet_name=s, header=None)
        if len(df) <= 3:
            print(f"  Sheet {s}: empty/too short ({len(df)} rows)")
            continue
        data = df.iloc[3:, :13].copy()
        # forward fill month (col 0)
        data.iloc[:, 0] = data.iloc[:, 0].ffill()
        # count valid rows: where school/location (col 2), topic (col 6), or instructor (col 7) is not null
        valid_rows = data[data.iloc[:, 2].notna() | data.iloc[:, 6].notna() | data.iloc[:, 7].notna()]
        print(f"  Sheet {s}: {len(valid_rows)} data rows")
        total_sessions += len(valid_rows)
    print(f"Total sessions for {academic_year}: {total_sessions}")

    # 2. School Enrollment sheets
    enrollment_sheets = [s for s in wb.sheetnames if 'school' in s.lower() and 'data' in s.lower()]
    print("Found enrollment sheets:", enrollment_sheets)
    total_enrollments = 0
    for s in enrollment_sheets:
        df = pd.read_excel(file_path, sheet_name=s, header=None)
        # Find row where column index 1 contains 'name of the school' or 'school'
        header_row_idx = None
        for idx in range(min(10, len(df))):
            val = str(df.iloc[idx, 1]).strip().lower()
            if 'name of the school' in val or 'name of school' in val or 'school' in val:
                header_row_idx = idx
                break
        if header_row_idx is None:
            print(f"  Sheet {s}: could not find 'Name of the School' header!")
            continue
        
        # Sub-header is at header_row_idx + 1 (B, G, T)
        # Data rows start at header_row_idx + 2
        data_rows = df.iloc[header_row_idx + 2:].copy()
        valid_schools = 0
        for _, r in data_rows.iterrows():
            school_name = str(r.iloc[1]).strip()
            if not school_name or school_name.lower() in ['nan', 'none', 'total', 'grand total']:
                continue
            valid_schools += 1
        print(f"  Sheet {s}: header at row {header_row_idx}, {valid_schools} school rows")
        total_enrollments += valid_schools

    print(f"Total school rows for {academic_year}: {total_enrollments}")

if __name__ == '__main__':
    test_parse('Documentation/YfS-Aequs all Programme Weekly Tracker_2025-26 (1).xlsx', '2025-26')
    test_parse('Documentation/YFS-Aequs all Programme Weekly Tracker_2026-27.xlsx', '2026-27')
