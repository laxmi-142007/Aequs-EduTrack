import csv
from io import BytesIO
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse
from django.contrib import messages
from django.db.models import Q, Sum, Count

from .models import (
    Volunteer,
    VolunteerActivity,
    EventParticipation,
)
from events.models import Event


def volunteer_list(request):
    q = request.GET.get("q", "").strip()
    status_filter = request.GET.get("status", "").strip()
    gender_filter = request.GET.get("gender", "").strip()

    volunteers = Volunteer.objects.all()

    # KPI counts across all volunteers
    total_volunteers = volunteers.count()
    active_count = volunteers.filter(status="ACTIVE").count()
    total_hours_aggregated = EventParticipation.objects.aggregate(Sum("hours_contributed"))["hours_contributed__sum"] or 0
    total_participations = EventParticipation.objects.count()

    if q:
        volunteers = volunteers.filter(
            Q(first_name__icontains=q) |
            Q(last_name__icontains=q) |
            Q(email__icontains=q) |
            Q(phone__icontains=q) |
            Q(skills__icontains=q) |
            Q(occupation__icontains=q) |
            Q(city__icontains=q)
        )

    if status_filter:
        volunteers = volunteers.filter(status=status_filter)

    if gender_filter:
        volunteers = volunteers.filter(gender=gender_filter)

    # Annotate with participation count and contributed hours
    volunteers = volunteers.annotate(
        total_hours=Sum("event_participations__hours_contributed"),
        event_count=Count("event_participations"),
    )

    return render(
        request,
        "volunteers/volunteer_list.html",
        {
            "volunteers": volunteers,
            "q": q,
            "status_filter": status_filter,
            "gender_filter": gender_filter,
            "status_choices": Volunteer.STATUS_CHOICES,
            "gender_choices": Volunteer.GENDER_CHOICES,
            "total_volunteers": total_volunteers,
            "active_count": active_count,
            "total_hours": round(total_hours_aggregated, 1),
            "total_participations": total_participations,
        },
    )


def volunteer_create(request):
    if request.method == "POST":
        volunteer = Volunteer.objects.create(
            first_name=request.POST.get("first_name", "").strip(),
            last_name=request.POST.get("last_name", "").strip(),
            email=request.POST.get("email", "").strip(),
            phone=request.POST.get("phone", "").strip(),
            gender=request.POST.get("gender", ""),
            date_of_birth=request.POST.get("date_of_birth") or None,
            address=request.POST.get("address", ""),
            city=request.POST.get("city", ""),
            state=request.POST.get("state", ""),
            pincode=request.POST.get("pincode", ""),
            occupation=request.POST.get("occupation", ""),
            qualification=request.POST.get("qualification", ""),
            skills=request.POST.get("skills", ""),
            status=request.POST.get("status", "ACTIVE"),
            remarks=request.POST.get("remarks", ""),
        )

        messages.success(request, f"Volunteer {volunteer.first_name} {volunteer.last_name} enrolled successfully.")
        return redirect("volunteers:detail", pk=volunteer.pk)

    return render(
        request,
        "volunteers/volunteer_form.html",
        {
            "gender_choices": Volunteer.GENDER_CHOICES,
            "status_choices": Volunteer.STATUS_CHOICES,
        },
    )


def volunteer_detail(request, pk):
    volunteer = get_object_or_404(Volunteer, pk=pk)

    activities = volunteer.activities.all()
    participations = volunteer.event_participations.select_related("event").all()

    total_hours = participations.aggregate(Sum("hours_contributed"))["hours_contributed__sum"] or 0
    assigned_event_ids = participations.values_list("event_id", flat=True)
    available_events = Event.objects.exclude(id__in=assigned_event_ids)

    return render(
        request,
        "volunteers/volunteer_detail.html",
        {
            "volunteer": volunteer,
            "activities": activities,
            "participations": participations,
            "total_hours": round(total_hours, 1),
            "events_count": participations.count(),
            "activities_count": activities.count(),
            "available_events": available_events,
            "participation_status_choices": EventParticipation.STATUS_CHOICES,
            "activity_status_choices": VolunteerActivity.STATUS_CHOICES,
        },
    )


def volunteer_edit(request, pk):
    volunteer = get_object_or_404(Volunteer, pk=pk)

    if request.method == "POST":
        volunteer.first_name = request.POST.get("first_name", volunteer.first_name).strip()
        volunteer.last_name = request.POST.get("last_name", volunteer.last_name).strip()
        volunteer.email = request.POST.get("email", volunteer.email).strip()
        volunteer.phone = request.POST.get("phone", volunteer.phone).strip()
        volunteer.gender = request.POST.get("gender", volunteer.gender)
        volunteer.date_of_birth = request.POST.get("date_of_birth") or None
        volunteer.address = request.POST.get("address", volunteer.address)
        volunteer.city = request.POST.get("city", volunteer.city)
        volunteer.state = request.POST.get("state", volunteer.state)
        volunteer.pincode = request.POST.get("pincode", volunteer.pincode)
        volunteer.occupation = request.POST.get("occupation", volunteer.occupation)
        volunteer.qualification = request.POST.get("qualification", volunteer.qualification)
        volunteer.skills = request.POST.get("skills", volunteer.skills)
        volunteer.status = request.POST.get("status", volunteer.status)
        volunteer.remarks = request.POST.get("remarks", volunteer.remarks)
        volunteer.save()

        messages.success(request, f"Profile for {volunteer.first_name} {volunteer.last_name} updated.")
        return redirect("volunteers:detail", pk=volunteer.pk)

    return render(
        request,
        "volunteers/volunteer_form.html",
        {
            "volunteer": volunteer,
            "is_edit": True,
            "gender_choices": Volunteer.GENDER_CHOICES,
            "status_choices": Volunteer.STATUS_CHOICES,
        },
    )


def volunteer_delete(request, pk):
    volunteer = get_object_or_404(Volunteer, pk=pk)
    if request.method == "POST":
        name = str(volunteer)
        volunteer.delete()
        messages.success(request, f"Volunteer {name} was removed from the roster.")
        return redirect("volunteers:list")

    return redirect("volunteers:detail", pk=volunteer.pk)


def activity_create(request, pk):
    volunteer = get_object_or_404(Volunteer, pk=pk)

    if request.method == "POST":
        VolunteerActivity.objects.create(
            volunteer=volunteer,
            activity_name=request.POST.get("activity_name"),
            activity_date=request.POST.get("activity_date"),
            description=request.POST.get("description", ""),
            status=request.POST.get("status", "ASSIGNED"),
            remarks=request.POST.get("remarks", ""),
        )

        messages.success(request, "Activity recorded successfully.")
        return redirect("volunteers:detail", pk=volunteer.pk)

    return render(
        request,
        "volunteers/activity_form.html",
        {
            "volunteer": volunteer,
            "status_choices": VolunteerActivity.STATUS_CHOICES,
        },
    )


def event_participation_create(request, pk):
    volunteer = get_object_or_404(Volunteer, pk=pk)
    events = Event.objects.all()

    if request.method == "POST":
        event_id = request.POST.get("event")
        role = request.POST.get("role", "").strip()
        status = request.POST.get("participation_status", "ASSIGNED")
        hours = request.POST.get("hours_contributed", 0)
        remarks = request.POST.get("remarks", "").strip()

        if event_id:
            event = get_object_or_404(Event, pk=event_id)
            try:
                hours_val = float(hours) if hours else 0
            except ValueError:
                hours_val = 0

            EventParticipation.objects.update_or_create(
                volunteer=volunteer,
                event=event,
                defaults={
                    "role": role,
                    "participation_status": status,
                    "hours_contributed": hours_val,
                    "remarks": remarks,
                },
            )
            messages.success(request, f"Assigned to event '{event.title}'.")
        return redirect("volunteers:detail", pk=volunteer.pk)

    return render(
        request,
        "volunteers/event_participation_form.html",
        {
            "volunteer": volunteer,
            "events": events,
            "status_choices": EventParticipation.STATUS_CHOICES,
        },
    )


def volunteer_export_csv(request):
    volunteers = Volunteer.objects.annotate(
        total_hours=Sum("event_participations__hours_contributed"),
        event_count=Count("event_participations"),
    ).all()

    q = request.GET.get("q", "").strip()
    status_filter = request.GET.get("status", "").strip()

    if q:
        volunteers = volunteers.filter(
            Q(first_name__icontains=q) |
            Q(last_name__icontains=q) |
            Q(email__icontains=q) |
            Q(phone__icontains=q) |
            Q(skills__icontains=q)
        )
    if status_filter:
        volunteers = volunteers.filter(status=status_filter)

    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="aequs_volunteers.csv"'

    writer = csv.writer(response)
    writer.writerow(["ID", "First Name", "Last Name", "Email", "Phone", "Gender", "Status", "Occupation", "Skills", "Events Attended", "Contributed Hours"])

    for v in volunteers:
        writer.writerow([
            v.id,
            v.first_name,
            v.last_name,
            v.email,
            v.phone,
            v.get_gender_display(),
            v.get_status_display(),
            v.occupation or "—",
            v.skills or "—",
            v.event_count,
            round(v.total_hours or 0, 1),
        ])

    return response


def volunteer_export_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    volunteers = Volunteer.objects.annotate(
        total_hours=Sum("event_participations__hours_contributed"),
        event_count=Count("event_participations"),
    ).all()

    q = request.GET.get("q", "").strip()
    status_filter = request.GET.get("status", "").strip()

    if q:
        volunteers = volunteers.filter(
            Q(first_name__icontains=q) |
            Q(last_name__icontains=q) |
            Q(email__icontains=q) |
            Q(phone__icontains=q) |
            Q(skills__icontains=q)
        )
    if status_filter:
        volunteers = volunteers.filter(status=status_filter)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Volunteers Roster"

    title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    meta_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="DD2D27", end_color="DD2D27", fill_type="solid")
    border_side = Side(border_style="thin", color="E2E8F0")
    cell_border = Border(top=border_side, bottom=border_side, left=border_side, right=border_side)

    ws.append(["Aequs Foundation - Volunteers Roster"])
    ws.cell(row=1, column=1).font = title_font
    ws.append([f"Total Enrolled: {len(volunteers)} volunteers | Status Filter: {status_filter or 'All'}"])
    ws.cell(row=2, column=1).font = meta_font
    ws.append([])

    headers = ["ID", "Name", "Email", "Phone", "Gender", "Status", "Occupation", "City", "Events", "Contributed Hours"]
    ws.append(headers)
    header_row_idx = ws.max_row
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=header_row_idx, column=col_idx)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center" if headers[col_idx-1] in ["ID", "Gender", "Status", "Events", "Contributed Hours"] else "left")

    for v in volunteers:
        full_name = f"{v.first_name} {v.last_name}".strip()
        ws.append([
            v.id,
            full_name,
            v.email,
            v.phone,
            v.get_gender_display(),
            v.get_status_display(),
            v.occupation or "—",
            v.city or "—",
            v.event_count,
            round(v.total_hours or 0, 1),
        ])
        curr_row = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=curr_row, column=col_idx)
            c.border = cell_border

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 35)

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)

    response = HttpResponse(
        bio.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="aequs_volunteers.xlsx"'
    return response