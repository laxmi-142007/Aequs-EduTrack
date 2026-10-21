import re
import os
import io
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from django.db import transaction
from django.utils.text import slugify
from schools.models import School
from programs.models import NGO
from trackers.models import Programme, SchoolEnrollment, WeeklyTrackerSession

# ==============================================================================
# CANONICALIZATION & ALIAS MAPPINGS
# ==============================================================================

CANONICAL_SCHOOL_ALIASES = {
    "uhps ymd": "UHPS Yamakanamaradi",
    "uemps ymd": "UEMPS Yamakanamaradi",
    "khpgs ymd": "KHPGS Yamakanamaradi",
    "khpgs yamkanamaradi": "KHPGS Yamakanamaradi",
    "khpbs ymd": "KHPBS Yamakanamaradi",
    "ghs hunoor mastiholi": "GHS Hunnur Masthiholi",
    "ghs hunnur masthiholi": "GHS Hunnur Masthiholi",
    "pms khps hunnur massthiholi": "GHS Hunnur Masthiholi",
    "pms khps hunnur mastiholi": "GHS Hunnur Masthiholi",
}

MONTH_MAP = {
    "jan": 1, "january": 1,
    "feb": 2, "february": 2,
    "mar": 3, "march": 3,
    "apr": 4, "april": 4,
    "may": 5,
    "jun": 6, "june": 6,
    "jul": 7, "july": 7,
    "aug": 8, "august": 8,
    "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10,
    "nov": 11, "november": 11,
    "dec": 12, "december": 12,
}

TRACKER_SHEET_MAPPING = {
    # Youth for Seva (YFS)
    "hhp_weekly tracker": Programme.CHAAP_JR,
    "sep_weekly tracker": Programme.SEP,
    "ignite_tracker": Programme.IGNITE_BGM,
    "ahap_weekly tracker": Programme.CHAAP_YUVA,
    "nmms_weekly tracker": Programme.NMMS,
    "ignite(hbl)_tracker": Programme.IGNITE_HBL,
    "others": Programme.OTHERS,
    "winter & summer camp": Programme.CAMP,
    "winter & summer camp ": Programme.CAMP,
    "sap": Programme.SAP,
    "chaap junior weekly tracker": Programme.CHAAP_JR,
    "chaap-yuva_weekly tracker": Programme.CHAAP_YUVA,

    # Agastya International Foundation
    "msl": Programme.MSL,
    "msl_tracker": Programme.MSL,
    "msl weekly tracker": Programme.MSL,
    "mobile science lab": Programme.MSL,
    "mobile science lab tracker": Programme.MSL,
    "msl 1": Programme.MSL,
    "msl 2": Programme.MSL,
    "avishkar": Programme.AVISHKAR,
    "avishkar_tracker": Programme.AVISHKAR,
    "avishkar science fair": Programme.AVISHKAR,
    "vasantha": Programme.VASANTHA,
    "vasantha_tracker": Programme.VASANTHA,
    "vasantha after-school": Programme.VASANTHA,
    "yuvacharya": Programme.YUVACHARYA,
    "yuvacharya_tracker": Programme.YUVACHARYA,

    # Pratham Foundation
    "opn": Programme.OPN,
    "opn_tracker": Programme.OPN,
    "one notebook": Programme.OPN,
    "one notebook program": Programme.OPN,
    "reading": Programme.READING,
    "reading camp": Programme.READING,
    "learning & reading camp": Programme.READING,
    "reading_tracker": Programme.READING,
}

ENROLLMENT_SHEET_MAPPING = {
    # 2025-26 names
    "hhp school data": Programme.CHAAP_JR,
    "sep_schools_data": Programme.SEP,
    "ignite(bgm)_schools_data": Programme.IGNITE_BGM,
    "ahap_schools_data": Programme.CHAAP_YUVA,
    "nmms_schools_data": Programme.NMMS,
    "ignite(hbl)_schools_data": Programme.IGNITE_HBL,

    # 2026-27 names
    "chaap junior school data": Programme.CHAAP_JR,
    "chaap-yuva_schools_data": Programme.CHAAP_YUVA,
}


def normalize_school_name(raw_name):
    """
    Cleans raw school string and resolves canonical spelling aliases.
    """
    if not raw_name or pd.isna(raw_name):
        return ""
    cleaned = str(raw_name).strip()
    lookup = cleaned.lower()
    return CANONICAL_SCHOOL_ALIASES.get(lookup, cleaned)


canonicalize_school_name = normalize_school_name


def is_likely_activity_or_meeting(text):
    """
    Determines if text in the school column represents a non-school activity, meeting, or venue.
    """
    if not text:
        return True
    t = text.lower()
    activity_keywords = [
        "meeting", "presentation", "powerpoint", "workshop", "webinar",
        "planning", "training", "orientation", "discussion", "revis",
        "aequs campus", "office", "holiday", "leave"
    ]
    return any(k in t for k in activity_keywords)


def get_or_create_school_cached(raw_name, cache):
    """
    Finds or creates a School record using canonicalization and a local in-memory cache
    to eliminate redundant DB queries.
    """
    if not raw_name or pd.isna(raw_name):
        return None

    canonical = normalize_school_name(raw_name)
    if not canonical:
        return None

    cache_key = canonical.lower()
    if cache_key in cache:
        return cache[cache_key]

    # 1. Exact case-insensitive match on School.name
    school = School.objects.filter(name__iexact=canonical).first()

    # 2. Match on code_aliases if not found
    if not school:
        school = School.objects.filter(code_aliases__icontains=canonical).first()

    # 3. Create if missing and not an activity entry
    if not school:
        if is_likely_activity_or_meeting(canonical):
            cache[cache_key] = None
            return None

        # Generate a temporary unique UDISE code
        safe_slug = slugify(canonical)[:12] or "sch"
        udise = f"YFS-{safe_slug}-{School.objects.count() + 1000}"
        school = School.objects.create(
            name=canonical,
            udise_code=udise,
            district="Belagavi",
            status=School.Status.ACTIVE,
            code_aliases=[raw_name.strip()] if raw_name.strip() != canonical else [],
        )
    else:
        # Append raw_name to code_aliases if new
        raw_clean = raw_name.strip()
        current_aliases = school.code_aliases or []
        if raw_clean and raw_clean.lower() != school.name.lower() and raw_clean not in current_aliases:
            current_aliases.append(raw_clean)
            school.code_aliases = current_aliases
            school.save(update_fields=["code_aliases"])

    cache[cache_key] = school
    return school


get_or_create_school_canonical = get_or_create_school_cached


def parse_academic_date(month_val, day_val, academic_year):
    """
    Converts (month, day, academic_year) into a real datetime.date using Indian academic cycle:
    June (6) to Dec (12) -> start_year
    Jan (1) to May (5) -> start_year + 1
    """
    if not academic_year:
        return None, "", None

    try:
        start_year = int(academic_year.split("-")[0].strip())
    except (ValueError, IndexError):
        start_year = 2025

    if isinstance(month_val, (datetime.date, datetime.datetime)):
        return month_val.date() if isinstance(month_val, datetime.datetime) else month_val, month_val.strftime("%B"), month_val.day

    month_str = str(month_val or "").strip()
    month_key = month_str.lower()
    month_num = MONTH_MAP.get(month_key)
    if not month_num:
        # Check prefix matching (e.g. 'Jan' in 'January')
        for k, v in MONTH_MAP.items():
            if month_key.startswith(k):
                month_num = v
                break

    day_num = None
    if day_val is not None and not pd.isna(day_val):
        try:
            day_num = int(float(str(day_val).strip()))
        except (ValueError, TypeError):
            day_num = None

    if not month_num or not day_num or day_num < 1 or day_num > 31:
        return None, month_str, day_num

    year = start_year if month_num >= 6 else start_year + 1
    try:
        resolved_date = datetime.date(year, month_num, day_num)
        return resolved_date, month_str, day_num
    except ValueError:
        return None, month_str, day_num


def parse_flexible_date(date_val=None, month_val=None, day_val=None, academic_year="2026-27"):
    """
    Attempts to extract (session_date, month_name, day_number) from either
    a full calendar date (datetime, date string) or (month, day) academic cycle values.
    """
    if isinstance(date_val, datetime.datetime):
        d = date_val.date()
        return d, d.strftime("%B"), d.day
    if isinstance(date_val, datetime.date):
        return date_val, date_val.strftime("%B"), date_val.day

    if date_val and not pd.isna(date_val):
        s = str(date_val).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d-%b-%Y", "%d/%m/%y", "%d-%m-%y", "%m/%d/%Y"):
            try:
                parsed = datetime.datetime.strptime(s.split()[0], fmt).date()
                return parsed, parsed.strftime("%B"), parsed.day
            except (ValueError, TypeError):
                continue

    if month_val is not None or day_val is not None:
        return parse_academic_date(month_val, day_val, academic_year)
    return None, "", None


def parse_safe_int(val, default=0):
    if val is None or pd.isna(val):
        return default
    try:
        return int(float(str(val).replace(",", "").strip()))
    except (ValueError, TypeError):
        return default


def parse_boolean_teacher(val):
    if val is None or pd.isna(val):
        return False
    s = str(val).strip().lower()
    return s in ["yes", "y", "true", "1", "present"]


# ==============================================================================
# MAIN WORKBOOK INGESTION SERVICE
# ==============================================================================

def import_yfs_workbook(file_obj_or_path, academic_year, ngo_code=None, target_program=None, target_programme=None):
    """
    Idempotent, transaction-safe ingestion service for Weekly Trackers and Session Logs.
    Supports:
      1. Multi-sheet partner NGO workbooks (YFS, Agastya, Pratham).
      2. Single-sheet Weekly Session Log files uploaded directly for a specific Program.
      3. Automatic resolution of target Programme and linking of partner schools.
      4. Automatic resolution of Partner NGO without requiring user selection.
    """
    # 1. Resolve Target Programme from Program if not explicitly given
    if target_program and not target_programme:
        from programs.views import resolve_tracker_programme
        target_programme = resolve_tracker_programme(target_program)

    # 2. Resolve default NGO from program or target_programme if available
    ngo = None
    if target_program:
        if target_program.ngo:
            ngo = target_program.ngo
        elif target_program.project and target_program.project.ngo:
            ngo = target_program.project.ngo

    ngo_names = {
        "AGASTYA": "Agastya International Foundation",
        "PRATHAM": "Pratham Foundation",
        "YFS": "Youth for Seva",
    }

    target_ngo_code_resolved = None
    if target_programme:
        if target_programme in [Programme.MSL, Programme.AVISHKAR, Programme.VASANTHA, Programme.YUVACHARYA]:
            target_ngo_code_resolved = "AGASTYA"
        elif target_programme in [Programme.OPN, Programme.READING, "CAMP"]:
            target_ngo_code_resolved = "PRATHAM"
        elif target_programme in [Programme.CHAAP_JR, Programme.CHAAP_YUVA, Programme.SEP, Programme.IGNITE_BGM, Programme.IGNITE_HBL, Programme.NMMS]:
            target_ngo_code_resolved = "YFS"

    if target_ngo_code_resolved and not ngo:
        ngo = NGO.objects.filter(code__iexact=target_ngo_code_resolved).first()
        if not ngo:
            ngo = NGO.objects.create(
                name=ngo_names.get(target_ngo_code_resolved, f"{target_ngo_code_resolved} Foundation"),
                code=target_ngo_code_resolved,
            )

    if not ngo and ngo_code:
        ngo = NGO.objects.filter(code__iexact=ngo_code).first()
        if not ngo:
            ngo = NGO.objects.create(
                name=ngo_names.get(ngo_code.upper(), f"{ngo_code.upper()} Foundation"),
                code=ngo_code.upper(),
            )

    excel_file = pd.ExcelFile(file_obj_or_path)
    sheet_names = excel_file.sheet_names

    # 3. If NGO still not resolved, auto-detect from workbook sheet names
    if not ngo:
        clean_sheets = [s.strip().lower() for s in sheet_names]
        detected_code = None
        if any(any(k in s for k in ["msl", "avishkar", "vasantha", "yuvacharya"]) for s in clean_sheets):
            detected_code = "AGASTYA"
        elif any(any(k in s for k in ["opn", "reading", "one notebook"]) for s in clean_sheets):
            detected_code = "PRATHAM"
        elif any(any(k in s for k in ["chaap", "sep", "ignite", "nmms", "hhp", "ahap"]) for s in clean_sheets):
            detected_code = "YFS"

        if detected_code:
            ngo = NGO.objects.filter(code__iexact=detected_code).first()
            if not ngo:
                ngo = NGO.objects.create(
                    name=ngo_names.get(detected_code, f"{detected_code} Foundation"),
                    code=detected_code,
                )

    # Fallback to Youth for Seva (primary NGO partner) or first active NGO
    if not ngo:
        ngo = NGO.objects.filter(code__iexact="YFS").first() or NGO.objects.first()

    if not ngo:
        fallback_code = (ngo_code or "YFS").upper()
        ngo = NGO.objects.create(name=f"{fallback_code} Foundation", code=fallback_code)

    # Normalize academic year string format
    ay_clean = academic_year.strip() if academic_year else "2026-27"
    if not re.match(r"^20\d{2}-\d{2}$", ay_clean):
        # e.g. 2025 -> 2025-26
        if ay_clean.isdigit() and len(ay_clean) == 4:
            ay_clean = f"{ay_clean}-{str(int(ay_clean[-2:]) + 1).zfill(2)}"
        elif not ay_clean:
            ay_clean = "2026-27"

    school_cache = {}
    sessions_to_create = []
    enrollments_to_create = []
    sheet_stats = {}

    generic_sheet_keywords = [
        "weekly session log", "weekly session logs", "weekly tracker",
        "weekly sessions", "session log", "session logs", "sessions",
        "session", "log", "sheet1", "tracker", "field sessions"
    ]

    # --------------------------------------------------------------------------
    # 1. PARSE WEEKLY TRACKER & SESSION LOG SHEETS
    # --------------------------------------------------------------------------
    for sheet_name in sheet_names:
        clean_key = sheet_name.strip().lower()
        prog_choice = TRACKER_SHEET_MAPPING.get(clean_key)

        # Fallback matching for single-sheet or generic weekly session log files
        if not prog_choice:
            if clean_key in generic_sheet_keywords or len(sheet_names) == 1:
                prog_choice = target_programme or Programme.SEP
            elif target_programme and any(k in clean_key for k in ["tracker", "session", "log", "weekly"]):
                prog_choice = target_programme

        if not prog_choice:
            continue

        sheet_ngo = ngo
        if not target_program:
            if prog_choice in [Programme.MSL, Programme.AVISHKAR, Programme.VASANTHA, Programme.YUVACHARYA]:
                ag_ngo = NGO.objects.filter(code__iexact="AGASTYA").first()
                if ag_ngo:
                    sheet_ngo = ag_ngo
            elif prog_choice in [Programme.OPN, Programme.READING]:
                pr_ngo = NGO.objects.filter(code__iexact="PRATHAM").first()
                if pr_ngo:
                    sheet_ngo = pr_ngo
            elif prog_choice in [Programme.CHAAP_JR, Programme.CHAAP_YUVA, Programme.SEP, Programme.IGNITE_BGM, Programme.IGNITE_HBL, Programme.NMMS]:
                yf_ngo = NGO.objects.filter(code__iexact="YFS").first()
                if yf_ngo:
                    sheet_ngo = yf_ngo

        df = pd.read_excel(excel_file, sheet_name=sheet_name, header=None)
        if len(df) <= 1:
            continue

        # Detect column layout: check first 5 rows for column header titles
        col_map = {}
        header_row_idx = None

        for r_idx in range(min(5, len(df))):
            row_vals = [str(v or "").strip().lower() for v in df.iloc[r_idx]]
            has_school = any("school" in v or "venue" in v or "place" in v or "center" in v for v in row_vals)
            has_topic = any("topic" in v or "activity" in v or "module" in v or "subject" in v for v in row_vals)
            has_month_or_date = any("month" in v or "date" in v for v in row_vals)

            if has_school and (has_topic or has_month_or_date):
                header_row_idx = r_idx
                for c_idx, cell_str in enumerate(row_vals):
                    if "date" in cell_str and "date" not in col_map:
                        col_map["date"] = c_idx
                    elif "month" in cell_str and "month" not in col_map:
                        col_map["month"] = c_idx
                    elif "day" in cell_str and "day" not in col_map:
                        col_map["day"] = c_idx
                    elif ("school" in cell_str or "venue" in cell_str or "place" in cell_str) and "school" not in col_map:
                        col_map["school"] = c_idx
                    elif ("std" in cell_str or "class" in cell_str or "grade" in cell_str) and "std" not in col_map:
                        col_map["std"] = c_idx
                    elif ("div" in cell_str or "section" in cell_str) and "div" not in col_map:
                        col_map["div"] = c_idx
                    elif ("timing" in cell_str or "time" in cell_str or "duration" in cell_str) and "timing" not in col_map:
                        col_map["timing"] = c_idx
                    elif ("topic" in cell_str or "activity" in cell_str or "module" in cell_str or "concept" in cell_str) and "topic" not in col_map:
                        col_map["topic"] = c_idx
                    elif ("instructor" in cell_str or "trainer" in cell_str or "faculty" in cell_str or "facilitator" in cell_str) and "instructor" not in col_map:
                        col_map["instructor"] = c_idx
                    elif ("type" in cell_str or "mode" in cell_str) and "type" not in col_map:
                        col_map["type"] = c_idx
                    elif ("boy" in cell_str) and "boys" not in col_map:
                        col_map["boys"] = c_idx
                    elif ("girl" in cell_str) and "girls" not in col_map:
                        col_map["girls"] = c_idx
                    elif ("total" in cell_str or "exposure" in cell_str or "strength" in cell_str) and "total" not in col_map:
                        col_map["total"] = c_idx
                    elif ("teacher" in cell_str or "present" in cell_str) and "teacher" not in col_map:
                        col_map["teacher"] = c_idx
                break

        sheet_session_count = 0

        # If standard 13-column tracker format
        if header_row_idx is None or ("school" not in col_map and len(df.columns) >= 12):
            start_row = 3 if len(df) > 3 else (header_row_idx + 1 if header_row_idx is not None else 1)
            data = df.iloc[start_row:, :13].copy()
            # Forward-fill month column (col 0)
            data.iloc[:, 0] = data.iloc[:, 0].ffill()

            for _, row in data.iterrows():
                c_school = row.iloc[2] if len(row) > 2 else None
                c_topic = row.iloc[6] if len(row) > 6 else None
                c_inst = row.iloc[7] if len(row) > 7 else None

                if pd.isna(c_school) and pd.isna(c_topic) and pd.isna(c_inst):
                    continue

                month_val = row.iloc[0]
                day_val = row.iloc[1] if len(row) > 1 else None
                raw_loc = str(c_school).strip() if pd.notna(c_school) else ""

                if raw_loc.lower() in ["total", "grand total", "nan"]:
                    continue

                session_date, month_name, day_num = parse_flexible_date(
                    month_val=month_val, day_val=day_val, academic_year=ay_clean
                )

                school_obj = get_or_create_school_cached(raw_loc, school_cache) if raw_loc else None
                std_val = str(row.iloc[3]).strip() if len(row) > 3 and pd.notna(row.iloc[3]) else ""
                div_val = str(row.iloc[4]).strip() if len(row) > 4 and pd.notna(row.iloc[4]) else ""
                timing_val = str(row.iloc[5]).strip() if len(row) > 5 and pd.notna(row.iloc[5]) else ""
                topic_val = str(c_topic).strip() if pd.notna(c_topic) else ""
                inst_val = str(c_inst).strip() if pd.notna(c_inst) else ""
                sess_type = str(row.iloc[8]).strip() if len(row) > 8 and pd.notna(row.iloc[8]) else "Practical"
                boys_cnt = parse_safe_int(row.iloc[9]) if len(row) > 9 else 0
                girls_cnt = parse_safe_int(row.iloc[10]) if len(row) > 10 else 0
                total_exp = parse_safe_int(row.iloc[11]) if len(row) > 11 else 0
                if total_exp == 0 and (boys_cnt + girls_cnt > 0):
                    total_exp = boys_cnt + girls_cnt
                t_present = parse_boolean_teacher(row.iloc[12]) if len(row) > 12 else False

                sessions_to_create.append(
                    WeeklyTrackerSession(
                        ngo=ngo,
                        academic_year=ay_clean,
                        programme=prog_choice,
                        session_date=session_date,
                        month_name=month_name,
                        day_number=day_num,
                        school=school_obj,
                        raw_location_or_activity=raw_loc,
                        standard=std_val if std_val.lower() != "nan" else "",
                        division=div_val if div_val.lower() != "nan" else "",
                        timing=timing_val if timing_val.lower() != "nan" else "",
                        topic=topic_val if topic_val.lower() != "nan" else "",
                        instructor_name=inst_val if inst_val.lower() != "nan" else "",
                        session_type=sess_type if sess_type.lower() != "nan" else "Practical",
                        boys_present=boys_cnt,
                        girls_present=girls_cnt,
                        total_exposures=total_exp,
                        teacher_present=t_present,
                    )
                )
                sheet_session_count += 1
        else:
            # Named column mapping
            start_row = header_row_idx + 1
            # Check if sub-header exists (e.g. Boys, Girls under Students Present)
            if start_row < len(df):
                sub_row = [str(v or "").strip().lower() for v in df.iloc[start_row]]
                if "boys" in sub_row or "girls" in sub_row:
                    for s_idx, s_val in enumerate(sub_row):
                        if "boy" in s_val and "boys" not in col_map:
                            col_map["boys"] = s_idx
                        elif "girl" in s_val and "girls" not in col_map:
                            col_map["girls"] = s_idx
                        elif "total" in s_val and "total" not in col_map:
                            col_map["total"] = s_idx
                    start_row += 1

            data = df.iloc[start_row:].copy()
            if "month" in col_map:
                data.iloc[:, col_map["month"]] = data.iloc[:, col_map["month"]].ffill()

            for _, row in data.iterrows():
                raw_loc = ""
                if "school" in col_map and col_map["school"] < len(row):
                    val = row.iloc[col_map["school"]]
                    raw_loc = str(val).strip() if pd.notna(val) else ""

                topic_val = ""
                if "topic" in col_map and col_map["topic"] < len(row):
                    val = row.iloc[col_map["topic"]]
                    topic_val = str(val).strip() if pd.notna(val) else ""

                inst_val = ""
                if "instructor" in col_map and col_map["instructor"] < len(row):
                    val = row.iloc[col_map["instructor"]]
                    inst_val = str(val).strip() if pd.notna(val) else ""

                if not raw_loc and not topic_val and not inst_val:
                    continue
                if raw_loc.lower() in ["total", "grand total", "nan"]:
                    continue

                date_val = row.iloc[col_map["date"]] if "date" in col_map and col_map["date"] < len(row) else None
                month_val = row.iloc[col_map["month"]] if "month" in col_map and col_map["month"] < len(row) else None
                day_val = row.iloc[col_map["day"]] if "day" in col_map and col_map["day"] < len(row) else None

                session_date, month_name, day_num = parse_flexible_date(
                    date_val=date_val, month_val=month_val, day_val=day_val, academic_year=ay_clean
                )

                school_obj = get_or_create_school_cached(raw_loc, school_cache) if raw_loc else None

                std_val = str(row.iloc[col_map["std"]]).strip() if "std" in col_map and col_map["std"] < len(row) and pd.notna(row.iloc[col_map["std"]]) else ""
                div_val = str(row.iloc[col_map["div"]]).strip() if "div" in col_map and col_map["div"] < len(row) and pd.notna(row.iloc[col_map["div"]]) else ""
                timing_val = str(row.iloc[col_map["timing"]]).strip() if "timing" in col_map and col_map["timing"] < len(row) and pd.notna(row.iloc[col_map["timing"]]) else ""
                sess_type = str(row.iloc[col_map["type"]]).strip() if "type" in col_map and col_map["type"] < len(row) and pd.notna(row.iloc[col_map["type"]]) else "Practical"

                boys_cnt = parse_safe_int(row.iloc[col_map["boys"]]) if "boys" in col_map and col_map["boys"] < len(row) else 0
                girls_cnt = parse_safe_int(row.iloc[col_map["girls"]]) if "girls" in col_map and col_map["girls"] < len(row) else 0
                total_exp = parse_safe_int(row.iloc[col_map["total"]]) if "total" in col_map and col_map["total"] < len(row) else 0
                if total_exp == 0 and (boys_cnt + girls_cnt > 0):
                    total_exp = boys_cnt + girls_cnt

                t_present = parse_boolean_teacher(row.iloc[col_map["teacher"]]) if "teacher" in col_map and col_map["teacher"] < len(row) else False

                sessions_to_create.append(
                    WeeklyTrackerSession(
                        ngo=sheet_ngo,
                        academic_year=ay_clean,
                        programme=prog_choice,
                        session_date=session_date,
                        month_name=month_name,
                        day_number=day_num,
                        school=school_obj,
                        raw_location_or_activity=raw_loc,
                        standard=std_val if std_val.lower() != "nan" else "",
                        division=div_val if div_val.lower() != "nan" else "",
                        timing=timing_val if timing_val.lower() != "nan" else "",
                        topic=topic_val if topic_val.lower() != "nan" else "",
                        instructor_name=inst_val if inst_val.lower() != "nan" else "",
                        session_type=sess_type if sess_type.lower() != "nan" else "Practical",
                        boys_present=boys_cnt,
                        girls_present=girls_cnt,
                        total_exposures=total_exp,
                        teacher_present=t_present,
                    )
                )
                sheet_session_count += 1

        sheet_stats[sheet_name] = f"{sheet_session_count} sessions"

    # --------------------------------------------------------------------------
    # 2. PARSE SCHOOL ENROLLMENT SHEETS (*_Schools_Data)
    # --------------------------------------------------------------------------
    seen_enrollment_keys = set()
    for sheet_name in sheet_names:
        clean_key = sheet_name.strip().lower()
        prog_choice = ENROLLMENT_SHEET_MAPPING.get(clean_key)
        if not prog_choice:
            continue

        df = pd.read_excel(excel_file, sheet_name=sheet_name, header=None)
        if len(df) <= 2:
            continue

        # Locate row with 'Name of the School' in col 1
        header_idx = None
        for r_idx in range(min(10, len(df))):
            val = str(df.iloc[r_idx, 1]).strip().lower()
            if "name of the school" in val or "name of school" in val or val == "school":
                header_idx = r_idx
                break

        if header_idx is None:
            continue

        header_row = list(df.iloc[header_idx])
        data_rows = df.iloc[header_idx + 2:].copy()

        # Identify standard triplets starting at col 2
        std_columns = []
        for col_idx in range(2, len(header_row)):
            cell_val = str(header_row[col_idx] or "").strip()
            # Match standards like '1st STD', '10th STD', '8th SDT', etc.
            if any(k in cell_val.lower() for k in ["std", "sdt", "class", "grade"]):
                std_clean = cell_val.replace("SDT", "STD").replace("sdt", "STD").strip()
                std_columns.append((col_idx, std_clean))

        sheet_enr_count = 0
        for _, r in data_rows.iterrows():
            raw_school = str(r.iloc[1]).strip() if pd.notna(r.iloc[1]) else ""
            if not raw_school or raw_school.lower() in ["nan", "none", "total", "grand total"]:
                continue

            school_obj = get_or_create_school_cached(raw_school, school_cache)
            if not school_obj:
                continue

            for col_idx, std_name in std_columns:
                boys_c = parse_safe_int(r.iloc[col_idx]) if col_idx < len(r) else 0
                girls_c = parse_safe_int(r.iloc[col_idx + 1]) if col_idx + 1 < len(r) else 0
                total_c = parse_safe_int(r.iloc[col_idx + 2]) if col_idx + 2 < len(r) else 0

                if total_c == 0 and (boys_c + girls_c > 0):
                    total_c = boys_c + girls_c

                enr_key = (ay_clean, prog_choice, school_obj.id, std_name)
                if enr_key in seen_enrollment_keys:
                    continue
                seen_enrollment_keys.add(enr_key)

                enrollments_to_create.append(
                    SchoolEnrollment(
                        ngo=ngo,
                        academic_year=ay_clean,
                        programme=prog_choice,
                        school=school_obj,
                        standard=std_name,
                        boys=boys_c,
                        girls=girls_c,
                        total=total_c,
                    )
                )
                sheet_enr_count += 1

        sheet_stats[sheet_name] = f"{sheet_enr_count} class enrollments"

    # --------------------------------------------------------------------------
    # 3. TRANSACTION-SAFE IDEMPOTENT PERSISTENCE
    # --------------------------------------------------------------------------
    with transaction.atomic():
        # Scoped deletion per (ngo, programme) pair
        prog_ngo_pairs = set((s.ngo, s.programme) for s in sessions_to_create)
        for s_ngo, s_prog in prog_ngo_pairs:
            del_filter = {"academic_year": ay_clean, "programme": s_prog}
            if s_ngo:
                del_filter["ngo"] = s_ngo
            WeeklyTrackerSession.objects.filter(**del_filter).delete()

        if not prog_ngo_pairs and target_programme:
            del_filter = {"academic_year": ay_clean, "programme": target_programme}
            if ngo:
                del_filter["ngo"] = ngo
            WeeklyTrackerSession.objects.filter(**del_filter).delete()

        enr_ngo_pairs = set((ngo, e.programme) for e in enrollments_to_create)
        for e_ngo, e_prog in enr_ngo_pairs:
            del_filter = {"academic_year": ay_clean, "programme": e_prog}
            if e_ngo:
                del_filter["ngo"] = e_ngo
            SchoolEnrollment.objects.filter(**del_filter).delete()

        WeeklyTrackerSession.objects.bulk_create(sessions_to_create, batch_size=500)
        SchoolEnrollment.objects.bulk_create(enrollments_to_create, batch_size=500)

        # Link matched schools to NGO
        valid_schools = [s for s in school_cache.values() if s]
        if ngo and valid_schools:
            ngo.partner_schools.add(*valid_schools)

        # Link matched schools to target Program Center if specified
        if target_program and valid_schools:
            target_program.participating_schools.add(*valid_schools)

    total_exposures = sum(s.total_exposures for s in sessions_to_create)
    return {
        "academic_year": ay_clean,
        "sessions_imported": len(sessions_to_create),
        "enrollments_imported": len(enrollments_to_create),
        "schools_count": len([s for s in school_cache.values() if s]),
        "total_exposures": total_exposures,
        "sheet_breakdown": sheet_stats,
    }


# ==============================================================================
# SAMPLE EXCEL TEMPLATE GENERATOR
# ==============================================================================

def generate_sample_weekly_session_log_excel(programme_title="Weekly Session Log", academic_year="2026-27"):
    """
    Generates a pre-formatted Microsoft Excel (.xlsx) template with standard weekly session
    tracking headers, subheadings, sample rows, and styled columns for fast user input.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly Session Log"

    # Title row
    ws.merge_cells("A1:M1")
    title_cell = ws["A1"]
    title_cell.value = f"Aequs Foundation — {programme_title} ({academic_year})"
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="166534", end_color="166534", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    headers = [
        "Month", "Day", "Name of the Schools / Place", "Std", "Division",
        "Timing", "Topic", "Instructor Name", "Session Type",
        "Boys", "Girls", "Total", "Teacher Present"
    ]
    ws.append(headers)
    ws.row_dimensions[2].height = 24

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1'),
    )

    for col_idx in range(1, 14):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    sample_rows = [
        ["July", 15, "GHS Hunnur Masthiholi", "Class 8", "A", "10:00 - 11:30 AM", "Scientific Method & Force Workshop", "Ramesh Patil", "Practical", 18, 22, 40, "Yes"],
        ["July", 22, "UHPS Yamakanamaradi", "Class 9", "B", "02:00 - 03:30 PM", "Solar Energy, Circuits & Optics", "Sunita Desai", "Theory", 15, 20, 35, "Yes"],
        ["August", 5, "KHPGS Yamakanamaradi", "Class 7", "A", "11:00 AM - 12:30 PM", "Robotics & Micro-controllers", "Anand Kulkarni", "Hands-on", 20, 25, 45, "Yes"],
    ]

    row_font = Font(name="Calibri", size=10)
    for r_idx, r_data in enumerate(sample_rows, start=3):
        ws.append(r_data)
        ws.row_dimensions[r_idx].height = 20
        for col_idx in range(1, 14):
            c = ws.cell(row=r_idx, column=col_idx)
            c.font = row_font
            c.border = thin_border
            if col_idx in (1, 2, 4, 5, 9, 10, 11, 12, 13):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

    col_widths = {1: 12, 2: 8, 3: 30, 4: 12, 5: 10, 6: 18, 7: 35, 8: 20, 9: 14, 10: 10, 11: 10, 12: 10, 13: 16}
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
