import json
import csv

from django.contrib import messages
from django.shortcuts import render, get_object_or_404, redirect
from django.urls import reverse
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_http_methods
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Sum, Count, Q

from .models import (
    Distribution,
    BenefitType,
    RecipientType,
    SchoolEssentialType,
    StudyKit,
    StudyKitItem,
)

from .forms import DistributionForm

from .services import (
    distribute_books,
    distribute_workbooks,
    issue_study_kit,
    issue_laptop_scholarship,
    distribute_school_essentials,
    get_top_puc_students,
    get_eligible_students_for_benefit,
)

from students.models import Student
from schools.models import School, SchoolResource, GradeStrength

from inventory.models import (
    InventoryItem,
    Laptop,
    InventoryCategory,
    StockTransaction,
)


# ============================================================
# HELPER
# ============================================================

def _distribution_to_dict(dist):
    return {
        "id": dist.id,
        "benefit_type": dist.benefit_type,
        "benefit_type_display": dist.get_benefit_type_display(),

        "recipient_type": dist.recipient_type,

        "recipient_name": (
            dist.student.student_name
            if dist.student
            else (
                dist.school.name
                if dist.school
                else "General"
            )
        ),

        "student_id": (
            dist.student.id
            if dist.student
            else None
        ),

        "student_name": (
            dist.student.student_name
            if dist.student
            else None
        ),

        "admission_number": (
            dist.student.admission_number
            if dist.student
            else ""
        ),

        "student_class": (
            dist.student.current_class
            if dist.student
            else ""
        ),

        "school_id": (
            dist.school.id
            if dist.school
            else (
                dist.student.school.id
                if dist.student and dist.student.school
                else None
            )
        ),

        "school_name": (
            dist.school.name
            if dist.school
            else (
                dist.student.school.name
                if dist.student and dist.student.school
                else "General"
            )
        ),

        "school_udise": (
            dist.school.udise_code
            if dist.school
            else (
                dist.student.school.udise_code
                if dist.student and dist.student.school
                else ""
            )
        ),

        "item_id": (
            dist.inventory_item.id
            if dist.inventory_item
            else None
        ),

        "item_name": (
            dist.inventory_item.item_name
            if dist.inventory_item
            else (
                dist.study_kit.name
                if dist.study_kit
                else "-"
            )
        ),

        "item_sku": (
            dist.inventory_item.sku
            if dist.inventory_item
            else ""
        ),

        "item_unit": (
            dist.inventory_item.unit
            if dist.inventory_item
            else "Units"
        ),

        "essential_type": dist.essential_item_type or "",

        "essential_type_display": (
            dist.get_essential_item_type_display()
            if dist.essential_item_type
            else ""
        ),

        "study_kit_name": (
            dist.study_kit.name
            if dist.study_kit
            else ""
        ),

        "quantity": dist.quantity,

        "academic_year": dist.academic_year,

        "distribution_date": (
            dist.distribution_date.strftime("%Y-%m-%d")
        ),

        "issued_by": dist.issued_by or "Admin",

        "remarks": dist.remarks or "",

        "created_at": (
            dist.created_at.strftime("%Y-%m-%d %H:%M")
        ),
    }


# ============================================================
# STOCK HELPERS
# ============================================================

def _stock_out(
    item,
    quantity,
    source_destination,
    performed_by,
    notes="",
):
    """
    Remove quantity from inventory and create STOCK_OUT transaction.
    """

    if not item or quantity <= 0:
        return

    item = InventoryItem.objects.select_for_update().get(
        pk=item.pk
    )

    if item.current_stock < quantity:
        raise ValidationError(
            f"Insufficient stock for '{item.item_name}'. "
            f"Available stock: {item.current_stock}. "
            f"Required: {quantity}."
        )

    previous_stock = item.current_stock

    item.current_stock -= quantity

    item.save(
        update_fields=[
            "current_stock",
            "updated_at",
        ]
    )

    StockTransaction.objects.create(
        item=item,
        transaction_type=StockTransaction.TransactionType.STOCK_OUT,
        quantity=quantity,
        previous_stock=previous_stock,
        new_stock=item.current_stock,
        source_destination=source_destination,
        performed_by=performed_by or "Admin",
        notes=notes,
    )


def _stock_in(
    item,
    quantity,
    source_destination,
    performed_by,
    notes="",
):
    """
    Return quantity to inventory and create STOCK_IN transaction.
    """

    if not item or quantity <= 0:
        return

    item = InventoryItem.objects.select_for_update().get(
        pk=item.pk
    )

    previous_stock = item.current_stock

    item.current_stock += quantity

    item.save(
        update_fields=[
            "current_stock",
            "updated_at",
        ]
    )

    StockTransaction.objects.create(
        item=item,
        transaction_type=StockTransaction.TransactionType.STOCK_IN,
        quantity=quantity,
        previous_stock=previous_stock,
        new_stock=item.current_stock,
        source_destination=source_destination,
        performed_by=performed_by or "Admin",
        notes=notes,
    )


def _get_distribution_roster_data():
    part_schools = list(
        School.objects.filter(grade_strengths__isnull=False)
        .distinct()
        .prefetch_related("grade_strengths", "resources")
    )
    if not part_schools:
        return None

    schools_output = []
    tot_c4 = [0, 0, 0]
    tot_c5 = [0, 0, 0]
    tot_c6 = [0, 0, 0]
    tot_c7 = [0, 0, 0]
    tot_c8 = [0, 0, 0]
    tot_c9 = [0, 0, 0]
    tot_c10 = [0, 0, 0]
    tot_gt = [0, 0, 0]

    for idx, sc in enumerate(part_schools, 1):
        teacher_res = sc.resources.filter(resource_name__icontains="teacher").first()
        teacher_name = teacher_res.resource_name.replace("Science Teacher:", "").strip() if teacher_res else ""
        teacher_mob = teacher_res.details.replace("Mobile:", "").strip() if teacher_res else ""

        grades_by_name = {g.grade_level.upper(): g for g in sc.grade_strengths.all()}
        def _get_g(name_key):
            for k, g in grades_by_name.items():
                if name_key in k:
                    return [g.male_students, g.female_students, g.total_students]
            return [0, 0, 0]

        c4 = _get_g("4")
        c5 = _get_g("5")
        c6 = _get_g("6")
        c7 = _get_g("7")
        c8 = _get_g("8")
        c9 = _get_g("9")
        c10 = _get_g("10")
        all_grades = [c4, c5, c6, c7, c8, c9, c10]
        gt_b = sum(g[0] for g in all_grades)
        gt_g = sum(g[1] for g in all_grades)
        gt_t = sc.student_strength or (gt_b + gt_g)
        gt = [gt_b, gt_g, gt_t]

        for i in range(3):
            tot_c4[i] += c4[i]
            tot_c5[i] += c5[i]
            tot_c6[i] += c6[i]
            tot_c7[i] += c7[i]
            tot_c8[i] += c8[i]
            tot_c9[i] += c9[i]
            tot_c10[i] += c10[i]
            tot_gt[i] += gt[i]

        schools_output.append({
            "sl": idx,
            "name": sc.name,
            "taluk": sc.taluk or sc.district or "Dharwad",
            "udise": sc.udise_code,
            "hm": sc.headmaster_name or "—",
            "hm_mob": sc.headmaster_phone or "—",
            "teacher": teacher_name or "—",
            "teacher_mob": teacher_mob or "—",
            "school_obj": sc,
            "rowspan": 1,
            "total_students": gt_t,
            "divisions": [{
                "division": "",
                "c4": c4, "c5": c5, "c6": c6, "c7": c7, "c8": c8, "c9": c9, "c10": c10,
                "gt": gt
            }]
        })

    taluks = sorted(list(set(s["taluk"] for s in schools_output if s["taluk"])))
    return {
        "title_org": "AGASTYA INTERNATIONAL FOUNDATION, NORTH KARNATAKA-1 REGION",
        "profile_title": "School Profile, Academic Year 2025-26",
        "program_title": "Mobile Science Lab / Study Kit Distributions",
        "donor_title": "Aequs Foundation Location: Dharwad",
        "academic_year": "2025-26",
        "schools": schools_output,
        "totals": {
            "c4": tot_c4, "c5": tot_c5, "c6": tot_c6, "c7": tot_c7,
            "c8": tot_c8, "c9": tot_c9, "c10": tot_c10, "gt": tot_gt
        },
        "taluks": taluks,
        "kpis": {
            "total_schools": len(schools_output),
            "total_students": tot_gt[2],
            "total_boys": tot_gt[0],
            "total_girls": tot_gt[1],
            "dlc_students": tot_c6[2],
            "taluks_count": len(taluks),
        }
    }


def _get_distribution_opn_data():
    opn_dists = list(
        Distribution.objects.filter(recipient_type=RecipientType.SCHOOL, remarks__icontains="Operation Notebook")
        .select_related("school")
        .order_by("location", "school__name")
    )
    if not opn_dists:
        return None

    seen_schools = set()
    schools_output = []
    tot_qty = 0
    idx = 1
    for d in opn_dists:
        if not d.school:
            continue
        s_key = (d.school.name.lower(), d.academic_year)
        if s_key in seen_schools:
            continue
        seen_schools.add(s_key)

        status = "Completed"
        if "Status: " in d.remarks:
            status = d.remarks.split("Status: ")[-1].strip()

        schools_output.append({
            "sl": idx,
            "name": d.school.name,
            "location": d.location or d.school.district or "Belagavi",
            "quantity": d.quantity,
            "issued_by": d.issued_by or "Aequs Foundation",
            "status": status,
            "academic_year": d.academic_year,
            "distribution_date": d.distribution_date.strftime("%Y-%m-%d") if d.distribution_date else "",
        })
        tot_qty += d.quantity
        idx += 1

    locations = sorted(list(set(s["location"] for s in schools_output if s["location"])))
    return {
        "title": "AEQUS FOUNDATION - OPERATION NOTEBOOK (OPN) DISTRIBUTION",
        "academic_year": "2026-27",
        "schools": schools_output,
        "total_notebooks": tot_qty,
        "total_schools": len(schools_output),
        "locations": locations,
    }


# ============================================================
# DISTRIBUTION LIST
# ============================================================

def distribution_list(request):
    """
    Display all distribution records with search, filter, and summary metrics,
    supporting both rich Excel Tabular View and Individual Records view.
    """
    q = request.GET.get("q", "").strip()
    benefit_filter = request.GET.get("benefit_type", "").strip()
    active_view = request.GET.get("view", "tabular" if not q and not benefit_filter else "records").lower()

    distributions = (
        Distribution.objects
        .select_related(
            "student",
            "student__school",
            "school",
            "inventory_item",
            "study_kit",
        )
        .order_by(
            "-distribution_date",
            "-created_at",
        )
    )

    if q:
        distributions = distributions.filter(
            Q(student__student_name__icontains=q)
            | Q(student__admission_number__icontains=q)
            | Q(school__name__icontains=q)
            | Q(location__icontains=q)
            | Q(issued_by__icontains=q)
            | Q(remarks__icontains=q)
        )

    if benefit_filter:
        distributions = distributions.filter(benefit_type=benefit_filter)

    total_distributions = distributions.count()

    individual_students = (
        distributions
        .filter(student__isnull=False)
        .values("student")
        .distinct()
        .count()
    )
    school_students = (
        distributions
        .filter(recipient_type=RecipientType.SCHOOL)
        .aggregate(total=Sum("quantity"))["total"] or 0
    )
    students_benefited = individual_students + school_students

    total_quantity = (
        distributions.aggregate(
            total=Sum("quantity")
        )["total"] or 0
    )

    # Breakdown by benefit type
    benefit_stats = {}
    for choice_key, choice_label in BenefitType.choices:
        count = Distribution.objects.filter(benefit_type=choice_key).count()
        benefit_stats[choice_label] = count

    roster_data = _get_distribution_roster_data()
    opn_data = _get_distribution_opn_data()

    active_sheet = request.GET.get("sheet", "").lower()
    if not active_sheet:
        active_sheet = request.session.get("last_distribution_sheet", "agastya" if roster_data else "opn")

    return render(
        request,
        "distributions/distribution_list.html",
        {
            "distributions": distributions,
            "total_distributions": total_distributions,
            "students_benefited": students_benefited,
            "total_quantity": total_quantity,
            "benefit_choices": BenefitType.choices,
            "selected_benefit": benefit_filter,
            "search_query": q,
            "benefit_stats": benefit_stats,
            "active_view": active_view,
            "active_sheet": active_sheet,
            "roster_data": roster_data,
            "opn_data": opn_data,
        },
    )


# ============================================================
# CREATE DISTRIBUTION
# ============================================================

def distribution_create(request):
    """
    Create distribution.

    Inventory:
        Distribution 10
        Inventory 100
        Result 90
    """

    if request.method == "POST":

        form = DistributionForm(request.POST, skip_eligibility=True)

        if form.is_valid():

            try:

                with transaction.atomic():

                    dist = form.save(commit=False)

                    # Automatically assign student's school
                    if (
                        dist.student
                        and dist.student.school
                        and not dist.school
                    ):
                        dist.school = dist.student.school

                    # Recipient type
                    if dist.student:
                        dist.recipient_type = RecipientType.STUDENT

                    elif dist.school:
                        dist.recipient_type = RecipientType.SCHOOL

                    issued_by = (
                        dist.issued_by
                        or (
                            request.user.username
                            if request.user.is_authenticated
                            else "Admin"
                        )
                    )

                    # Inventory deduction
                    if (
                        dist.inventory_item
                        and dist.quantity
                        and dist.quantity > 0
                    ):

                        recipient_desc = (
                            f"{dist.student.student_name} "
                            f"({dist.student.current_class})"
                            if dist.student
                            else (
                                f"School: {dist.school.name}"
                                if dist.school
                                else "General"
                            )
                        )

                        _stock_out(
                            item=dist.inventory_item,
                            quantity=dist.quantity,
                            source_destination=recipient_desc,
                            performed_by=issued_by,
                            notes=(
                                f"Distribution created. "
                                f"{dist.remarks or ''}"
                            ),
                        )

                    dist.save()

                return redirect("distributions:list")

            except ValidationError as e:

                form.add_error(
                    "quantity",
                    str(e),
                )

        # Invalid form falls through

    else:
        form = DistributionForm()

    return render(
        request,
        "distributions/distribution_form.html",
        {
            "form": form,
        },
    )


# ============================================================
# EDIT DISTRIBUTION
# ============================================================

def distribution_edit(request, pk):
    """
    Edit an existing distribution.

    Same item:

        Old = 10
        New = 15
        Inventory -5

        Old = 15
        New = 8
        Inventory +7

    Different item:

        Old item + old quantity
        New item - new quantity
    """

    distribution = get_object_or_404(
        Distribution.objects.select_related(
            "student",
            "school",
            "inventory_item",
            "study_kit",
        ),
        pk=pk,
    )

    if request.method == "POST":

        old_quantity = distribution.quantity
        old_item_id = distribution.inventory_item_id

        form = DistributionForm(
            request.POST,
            instance=distribution,
        )

        if form.is_valid():

            try:

                with transaction.atomic():

                    # Save old item reference
                    old_item = None

                    if old_item_id:
                        old_item = (
                            InventoryItem.objects
                            .select_for_update()
                            .get(pk=old_item_id)
                        )

                    # Get new data
                    dist = form.save(commit=False)

                    # Automatically assign school
                    if (
                        dist.student
                        and dist.student.school
                        and not dist.school
                    ):
                        dist.school = dist.student.school

                    # Recipient type
                    if dist.student:
                        dist.recipient_type = RecipientType.STUDENT

                    elif dist.school:
                        dist.recipient_type = RecipientType.SCHOOL

                    issued_by = (
                        dist.issued_by
                        or (
                            request.user.username
                            if request.user.is_authenticated
                            else "Admin"
                        )
                    )

                    new_item = None

                    if dist.inventory_item_id:

                        new_item = (
                            InventoryItem.objects
                            .select_for_update()
                            .get(
                                pk=dist.inventory_item_id
                            )
                        )

                    # ==================================================
                    # SAME ITEM
                    # ==================================================

                    if old_item_id == dist.inventory_item_id:

                        if new_item:

                            difference = (
                                dist.quantity - old_quantity
                            )

                            # ------------------------------------------
                            # NEW QUANTITY IS GREATER
                            # ------------------------------------------

                            if difference > 0:

                                try:

                                    _stock_out(
                                        item=new_item,
                                        quantity=difference,
                                        source_destination=(
                                            f"Edit Distribution "
                                            f"#{distribution.id}"
                                        ),
                                        performed_by=issued_by,
                                        notes=(
                                            f"Quantity increased "
                                            f"from {old_quantity} "
                                            f"to {dist.quantity}."
                                        ),
                                    )

                                except ValidationError as e:

                                    form.add_error(
                                        "quantity",
                                        str(e),
                                    )

                                    return render(
                                        request,
                                        "distributions/distribution_form.html",
                                        {
                                            "form": form,
                                            "distribution": distribution,
                                            "is_edit": True,
                                        },
                                    )

                            # ------------------------------------------
                            # NEW QUANTITY IS SMALLER
                            # ------------------------------------------

                            elif difference < 0:

                                returned_quantity = abs(
                                    difference
                                )

                                _stock_in(
                                    item=new_item,
                                    quantity=returned_quantity,
                                    source_destination=(
                                        f"Edit Distribution "
                                        f"#{distribution.id}"
                                    ),
                                    performed_by=issued_by,
                                    notes=(
                                        f"Quantity decreased "
                                        f"from {old_quantity} "
                                        f"to {dist.quantity}. "
                                        f"Returned "
                                        f"{returned_quantity} "
                                        f"to inventory."
                                    ),
                                )

                    # ==================================================
                    # INVENTORY ITEM CHANGED
                    # ==================================================

                    else:

                        # ------------------------------------------
                        # RETURN OLD ITEM
                        # ------------------------------------------

                        if old_item:

                            _stock_in(
                                item=old_item,
                                quantity=old_quantity,
                                source_destination=(
                                    f"Edit Distribution "
                                    f"#{distribution.id}"
                                ),
                                performed_by=issued_by,
                                notes=(
                                    "Old inventory item returned "
                                    "because distribution item "
                                    "was changed."
                                ),
                            )

                        # ------------------------------------------
                        # REMOVE NEW ITEM
                        # ------------------------------------------

                        if new_item and dist.quantity > 0:

                            try:

                                _stock_out(
                                    item=new_item,
                                    quantity=dist.quantity,
                                    source_destination=(
                                        f"Edit Distribution "
                                        f"#{distribution.id}"
                                    ),
                                    performed_by=issued_by,
                                    notes=(
                                        "New inventory item "
                                        "deducted because "
                                        "distribution item "
                                        "was changed."
                                    ),
                                )

                            except ValidationError as e:

                                form.add_error(
                                    "quantity",
                                    str(e),
                                )

                                return render(
                                    request,
                                    "distributions/distribution_form.html",
                                    {
                                        "form": form,
                                        "distribution": distribution,
                                        "is_edit": True,
                                    },
                                )

                    # Save edited distribution
                    dist.save()

                return redirect("distributions:list")

            except ValidationError as e:

                form.add_error(
                    "quantity",
                    str(e),
                )

    else:

        form = DistributionForm(
            instance=distribution,
        )

    return render(
        request,
        "distributions/distribution_form.html",
        {
            "form": form,
            "distribution": distribution,
            "is_edit": True,
        },
    )


# ============================================================
# DELETE DISTRIBUTION
# ============================================================

@require_http_methods(["POST"])
def distribution_delete(request, pk):
    """
    Delete distribution.

    Example:

        Inventory = 90
        Distribution = 10

        Delete distribution

        Inventory = 100
    """

    with transaction.atomic():

        distribution = get_object_or_404(
            Distribution.objects.select_related(
                "inventory_item",
            ),
            pk=pk,
        )

        quantity = distribution.quantity

        item = None

        if distribution.inventory_item_id:

            item = (
                InventoryItem.objects
                .select_for_update()
                .get(
                    pk=distribution.inventory_item_id
                )
            )

        issued_by = (
            request.user.username
            if request.user.is_authenticated
            else "Admin"
        )

        # Return quantity to inventory
        if item and quantity > 0:

            _stock_in(
                item=item,
                quantity=quantity,
                source_destination=(
                    f"Deleted Distribution "
                    f"#{distribution.id}"
                ),
                performed_by=issued_by,
                notes=(
                    f"Distribution deleted. "
                    f"Returned {quantity} "
                    f"item(s) to inventory."
                ),
            )

        # Delete distribution only after stock is returned
        distribution.delete()

    return redirect("distributions:list")


# ============================================================
# API - ELIGIBLE STUDENTS
# ============================================================

def api_eligible_students(request):

    benefit = request.GET.get(
        "benefit_type",
        "BOOK",
    ).upper()

    academic_year = request.GET.get(
        "academic_year",
        "",
    ).strip()

    students = get_eligible_students_for_benefit(
        benefit,
        academic_year=academic_year or None,
    )

    data = [
        {
            "id": s.id,
            "name": s.student_name,
            "admission_number": s.admission_number,
            "current_class": s.current_class,
            "school_id": (
                s.school.id
                if s.school
                else None
            ),
            "school_name": (
                s.school.name
                if s.school
                else "General"
            ),
        }
        for s in students
    ]

    return JsonResponse(
        {
            "success": True,
            "students": data,
            "count": len(data),
        }
    )


# ============================================================
# API - TOP PUC STUDENTS
# ============================================================

def api_top_puc_students(request):

    academic_year = request.GET.get(
        "academic_year"
    )

    top_students = get_top_puc_students(
        academic_year=academic_year,
        limit=10,
    )

    return JsonResponse(
        {
            "success": True,
            "top_students": top_students,
            "count": len(top_students),
        }
    )


# ============================================================
# API - STUDY KITS
# ============================================================

def api_study_kits(request):

    if request.method == "POST":

        try:

            if request.content_type == "application/json":
                data = json.loads(request.body)

            else:
                data = request.POST

            name = data.get(
                "name",
                "",
            ).strip()

            target_grade = data.get(
                "target_grade_level",
                "Class 1 to 2nd PUC",
            ).strip()

            description = data.get(
                "description",
                "",
            ).strip()

            items_list = data.get(
                "items",
                [],
            )

            if not name:

                return JsonResponse(
                    {
                        "success": False,
                        "error": "Kit Name is required.",
                    },
                    status=400,
                )

            kit = StudyKit.objects.create(
                name=name,
                target_grade_level=target_grade,
                description=description,
            )

            for it in items_list:

                if isinstance(it, dict):

                    it_name = (
                        it.get(
                            "item_name",
                            "",
                        ).strip()
                    )

                    qty = int(
                        it.get(
                            "quantity",
                            1,
                        )
                    )

                else:

                    it_name = str(it).strip()
                    qty = 1

                if it_name:

                    StudyKitItem.objects.create(
                        kit=kit,
                        item_name=it_name,
                        quantity_per_kit=qty,
                    )

            return JsonResponse(
                {
                    "success": True,
                    "message": (
                        f"Study Kit '{kit.name}' "
                        f"created successfully!"
                    ),
                }
            )

        except Exception as e:

            return JsonResponse(
                {
                    "success": False,
                    "error": str(e),
                },
                status=500,
            )

    kits = (
        StudyKit.objects
        .prefetch_related("items")
        .filter(is_active=True)
    )

    data = [

        {
            "id": k.id,
            "name": k.name,
            "target_grade_level": (
                k.target_grade_level
            ),
            "description": k.description,

            "items": [

                {
                    "item_name": it.item_name,
                    "quantity": it.quantity_per_kit,
                    "spec": it.specification,
                }

                for it in k.items.all()
            ],
        }

        for k in kits
    ]

    return JsonResponse(
        {
            "success": True,
            "kits": data,
            "count": len(data),
        }
    )


# ============================================================
# API - DISTRIBUTE BOOK
# ============================================================

@require_http_methods(["POST"])
def api_distribute_book(request):

    try:

        if request.content_type == "application/json":
            data = json.loads(request.body)

        else:
            data = request.POST

        student_id = data.get("student_id")
        item_id = data.get("item_id")

        quantity = int(
            data.get("quantity") or 1
        )

        academic_year = data.get(
            "academic_year",
            "2026-27",
        ).strip()

        remarks = data.get(
            "remarks",
            "",
        ).strip()

        if not student_id or not item_id:

            return JsonResponse(
                {
                    "success": False,
                    "error": (
                        "Student and Book item "
                        "are required."
                    ),
                },
                status=400,
            )

        student = get_object_or_404(
            Student,
            id=student_id,
        )

        item = get_object_or_404(
            InventoryItem,
            id=item_id,
        )

        issued_by = (
            request.user.username
            if request.user.is_authenticated
            else "Admin"
        )

        dist = distribute_books(
            student=student,
            item=item,
            quantity=quantity,
            academic_year=academic_year,
            issued_by=issued_by,
            remarks=remarks,
        )

        return JsonResponse(
            {
                "success": True,
                "message": (
                    f"Successfully distributed "
                    f"{quantity} copy of "
                    f"'{item.item_name}' to "
                    f"{student.student_name}. "
                    f"Inventory updated."
                ),
                "distribution": (
                    _distribution_to_dict(dist)
                ),
            }
        )

    except ValidationError as ve:

        return JsonResponse(
            {
                "success": False,
                "error": str(ve),
            },
            status=400,
        )

    except Exception as e:

        return JsonResponse(
            {
                "success": False,
                "error": str(e),
            },
            status=500,
        )


# ============================================================
# API - DISTRIBUTE WORKBOOK
# ============================================================

@require_http_methods(["POST"])
def api_distribute_workbook(request):

    try:

        if request.content_type == "application/json":
            data = json.loads(request.body)

        else:
            data = request.POST

        student_id = data.get("student_id")
        item_id = data.get("item_id")

        quantity = int(
            data.get("quantity") or 1
        )

        academic_year = data.get(
            "academic_year",
            "2026-27",
        ).strip()

        remarks = data.get(
            "remarks",
            "",
        ).strip()

        if not student_id or not item_id:

            return JsonResponse(
                {
                    "success": False,
                    "error": (
                        "Student and Workbook "
                        "item are required."
                    ),
                },
                status=400,
            )

        student = get_object_or_404(
            Student,
            id=student_id,
        )

        item = get_object_or_404(
            InventoryItem,
            id=item_id,
        )

        issued_by = (
            request.user.username
            if request.user.is_authenticated
            else "Admin"
        )

        dist = distribute_workbooks(
            student=student,
            item=item,
            quantity=quantity,
            academic_year=academic_year,
            issued_by=issued_by,
            remarks=remarks,
        )

        return JsonResponse(
            {
                "success": True,
                "message": (
                    f"Successfully distributed "
                    f"{quantity} workbook "
                    f"'{item.item_name}' to "
                    f"{student.student_name}. "
                    f"Inventory updated."
                ),
                "distribution": (
                    _distribution_to_dict(dist)
                ),
            }
        )

    except ValidationError as ve:

        return JsonResponse(
            {
                "success": False,
                "error": str(ve),
            },
            status=400,
        )

    except Exception as e:

        return JsonResponse(
            {
                "success": False,
                "error": str(e),
            },
            status=500,
        )


# ============================================================
# API - STUDY KIT
# ============================================================

@require_http_methods(["POST"])
def api_issue_study_kit(request):

    try:

        if request.content_type == "application/json":
            data = json.loads(request.body)

        else:
            data = request.POST

        student_id = data.get("student_id")
        item_id = data.get("item_id")
        kit_id = data.get("kit_id")

        quantity = int(
            data.get("quantity") or 1
        )

        academic_year = data.get(
            "academic_year",
            "2026-27",
        ).strip()

        remarks = data.get(
            "remarks",
            "",
        ).strip()

        if not student_id:

            return JsonResponse(
                {
                    "success": False,
                    "error": "Student is required.",
                },
                status=400,
            )

        student = get_object_or_404(
            Student,
            id=student_id,
        )

        item = (
            InventoryItem.objects
            .filter(id=item_id)
            .first()
            if item_id
            else None
        )

        kit = (
            StudyKit.objects
            .filter(id=kit_id)
            .first()
            if kit_id
            else None
        )

        issued_by = (
            request.user.username
            if request.user.is_authenticated
            else "Admin"
        )

        dist = issue_study_kit(
            student=student,
            item=item,
            kit=kit,
            quantity=quantity,
            academic_year=academic_year,
            issued_by=issued_by,
            remarks=remarks,
        )

        return JsonResponse(
            {
                "success": True,
                "message": (
                    f"Successfully issued Study Kit "
                    f"to {student.student_name}. "
                    f"Inventory updated."
                ),
                "distribution": (
                    _distribution_to_dict(dist)
                ),
            }
        )

    except ValidationError as ve:

        return JsonResponse(
            {
                "success": False,
                "error": str(ve),
            },
            status=400,
        )

    except Exception as e:

        return JsonResponse(
            {
                "success": False,
                "error": str(e),
            },
            status=500,
        )


# ============================================================
# API - LAPTOP SCHOLARSHIP
# ============================================================

@require_http_methods(["POST"])
def api_issue_laptop_scholarship(request):

    try:

        if request.content_type == "application/json":
            data = json.loads(request.body)

        else:
            data = request.POST

        student_id = data.get("student_id")
        laptop_id = data.get("laptop_id")

        academic_year = data.get(
            "academic_year",
            "2026-27",
        ).strip()

        notes = data.get(
            "notes",
            "",
        ).strip()

        if not student_id or not laptop_id:

            return JsonResponse(
                {
                    "success": False,
                    "error": (
                        "Student and Laptop "
                        "asset are required."
                    ),
                },
                status=400,
            )

        student = get_object_or_404(
            Student,
            id=student_id,
        )

        laptop = get_object_or_404(
            Laptop,
            id=laptop_id,
        )

        issued_by = (
            request.user.username
            if request.user.is_authenticated
            else "Admin"
        )

        assignment, dist = issue_laptop_scholarship(
            student=student,
            laptop=laptop,
            academic_year=academic_year,
            notes=notes,
            issued_by=issued_by,
        )

        return JsonResponse(
            {
                "success": True,
                "message": (
                    f"Laptop {laptop.asset_number} "
                    f"awarded to "
                    f"{student.student_name}."
                ),
                "distribution": (
                    _distribution_to_dict(dist)
                ),
            }
        )

    except ValidationError as ve:

        return JsonResponse(
            {
                "success": False,
                "error": str(ve),
            },
            status=400,
        )

    except Exception as e:

        return JsonResponse(
            {
                "success": False,
                "error": str(e),
            },
            status=500,
        )


# ============================================================
# API - SCHOOL ESSENTIALS
# ============================================================

@require_http_methods(["POST"])
def api_distribute_school_essentials(request):

    try:

        if request.content_type == "application/json":
            data = json.loads(request.body)

        else:
            data = request.POST

        school_id = data.get("school_id")
        item_id = data.get("item_id")

        essential_type = data.get(
            "essential_type",
            "OTHER",
        ).strip()

        quantity = int(
            data.get("quantity") or 1
        )

        academic_year = data.get(
            "academic_year",
            "2026-27",
        ).strip()

        remarks = data.get(
            "remarks",
            "",
        ).strip()

        if not school_id or not item_id:

            return JsonResponse(
                {
                    "success": False,
                    "error": (
                        "School and Essential "
                        "inventory item are required."
                    ),
                },
                status=400,
            )

        school = get_object_or_404(
            School,
            id=school_id,
        )

        item = get_object_or_404(
            InventoryItem,
            id=item_id,
        )

        issued_by = (
            request.user.username
            if request.user.is_authenticated
            else "Admin"
        )

        dist = distribute_school_essentials(
            school=school,
            item=item,
            essential_type=essential_type,
            quantity=quantity,
            academic_year=academic_year,
            issued_by=issued_by,
            remarks=remarks,
        )

        return JsonResponse(
            {
                "success": True,
                "message": (
                    f"Successfully allocated "
                    f"{quantity} {item.unit} of "
                    f"'{item.item_name}' to "
                    f"{school.name}. "
                    f"Inventory updated."
                ),
                "distribution": (
                    _distribution_to_dict(dist)
                ),
            }
        )

    except ValidationError as ve:

        return JsonResponse(
            {
                "success": False,
                "error": str(ve),
            },
            status=400,
        )

    except Exception as e:

        return JsonResponse(
            {
                "success": False,
                "error": str(e),
            },
            status=500,
        )


# ============================================================
# API - DISTRIBUTION HISTORY
# ============================================================

def api_distribution_history(request):

    benefit = request.GET.get(
        "benefit_type"
    )

    school_id = request.GET.get(
        "school_id"
    )

    student_id = request.GET.get(
        "student_id"
    )

    search = request.GET.get(
        "search",
        "",
    ).strip()

    qs = (
        Distribution.objects
        .select_related(
            "student",
            "school",
            "inventory_item",
            "study_kit",
        )
        .all()
    )

    if benefit:
        qs = qs.filter(
            benefit_type=benefit
        )

    if school_id:
        qs = qs.filter(
            Q(school_id=school_id)
            |
            Q(student__school_id=school_id)
        )

    if student_id:
        qs = qs.filter(
            student_id=student_id
        )

    if search:

        qs = qs.filter(
            Q(
                student__student_name__icontains=search
            )
            |
            Q(
                student__admission_number__icontains=search
            )
            |
            Q(
                school__name__icontains=search
            )
            |
            Q(
                inventory_item__item_name__icontains=search
            )
            |
            Q(
                remarks__icontains=search
            )
        )

    data = [
        _distribution_to_dict(d)
        for d in qs[:100]
    ]

    return JsonResponse(
        {
            "success": True,
            "distributions": data,
            "count": len(data),
        }
    )


# ============================================================
# API - SCHOOL REPORT
# ============================================================

def api_school_wise_report(request):

    schools = (
        School.objects
        .all()
        .order_by("name")
    )

    report = []

    for sc in schools:

        dists = Distribution.objects.filter(
            Q(school=sc)
            |
            Q(student__school=sc)
        )

        books_qty = (
            dists
            .filter(
                benefit_type=BenefitType.BOOK
            )
            .aggregate(total=Sum("quantity"))
            ["total"]
            or 0
        )

        workbooks_qty = (
            dists
            .filter(
                benefit_type=BenefitType.WORKBOOK
            )
            .aggregate(total=Sum("quantity"))
            ["total"]
            or 0
        )

        kits_qty = (
            dists
            .filter(
                benefit_type=BenefitType.STUDY_KIT
            )
            .aggregate(total=Sum("quantity"))
            ["total"]
            or 0
        )

        laptops_qty = (
            dists
            .filter(
                benefit_type=BenefitType.LAPTOP
            )
            .aggregate(total=Sum("quantity"))
            ["total"]
            or 0
        )

        essentials_qty = (
            dists
            .filter(
                benefit_type=BenefitType.SCHOOL_ESSENTIAL
            )
            .aggregate(total=Sum("quantity"))
            ["total"]
            or 0
        )

        report.append(
            {
                "school_id": sc.id,
                "school_name": sc.name,
                "udise_code": sc.udise_code,
                "district": sc.district,
                "books_distributed": books_qty,
                "workbooks_distributed": workbooks_qty,
                "study_kits_distributed": kits_qty,
                "laptops_issued": laptops_qty,
                "school_essentials_units": essentials_qty,
                "total_items_delivered": (
                    books_qty
                    + workbooks_qty
                    + kits_qty
                    + laptops_qty
                    + essentials_qty
                ),
            }
        )

    return JsonResponse(
        {
            "success": True,
            "report": report,
            "count": len(report),
        }
    )


# ============================================================
# API - STUDENT REPORT
# ============================================================

def api_student_wise_report(request):

    query = request.GET.get(
        "search",
        "",
    ).strip()

    students = (
        Student.objects
        .select_related("school")
        .all()
        .order_by("student_name")
    )

    if query:

        students = students.filter(
            Q(
                student_name__icontains=query
            )
            |
            Q(
                admission_number__icontains=query
            )
            |
            Q(
                school__name__icontains=query
            )
        )

    report = []

    for st in students:

        dists = Distribution.objects.filter(
            student=st
        )

        books_qty = (
            dists
            .filter(
                benefit_type=BenefitType.BOOK
            )
            .aggregate(total=Sum("quantity"))
            ["total"]
            or 0
        )

        workbooks_qty = (
            dists
            .filter(
                benefit_type=BenefitType.WORKBOOK
            )
            .aggregate(total=Sum("quantity"))
            ["total"]
            or 0
        )

        kits_qty = (
            dists
            .filter(
                benefit_type=BenefitType.STUDY_KIT
            )
            .aggregate(total=Sum("quantity"))
            ["total"]
            or 0
        )

        laptops_qty = (
            dists
            .filter(
                benefit_type=BenefitType.LAPTOP
            )
            .aggregate(total=Sum("quantity"))
            ["total"]
            or 0
        )

        report.append(
            {
                "student_id": st.id,
                "student_name": st.student_name,
                "admission_number": st.admission_number,
                "current_class": st.current_class,
                "school_name": (
                    st.school.name
                    if st.school
                    else "General"
                ),
                "books_received": books_qty,
                "workbooks_received": workbooks_qty,
                "study_kits_received": kits_qty,
                "laptops_received": laptops_qty,
                "total_benefits": (
                    books_qty
                    + workbooks_qty
                    + kits_qty
                    + laptops_qty
                ),
            }
        )

    return JsonResponse(
        {
            "success": True,
            "report": report,
            "count": len(report),
        }
    )


# ============================================================
# API - CSV EXPORT
# ============================================================

def api_export_distribution_csv(request):

    report_type = request.GET.get(
        "type",
        "school_wise",
    )

    response = HttpResponse(
        content_type="text/csv"
    )

    # --------------------------------------------------------
    # STUDENT WISE
    # --------------------------------------------------------

    if report_type == "student_wise":

        response[
            "Content-Disposition"
        ] = (
            'attachment; '
            'filename="student_distribution_report.csv"'
        )

        writer = csv.writer(response)

        writer.writerow(
            [
                "Student Name",
                "Admission No",
                "Class",
                "School",
                "Books",
                "Workbooks",
                "Study Kits",
                "Laptops",
                "Total Items",
            ]
        )

        for st in (
            Student.objects
            .select_related("school")
            .all()
            .order_by("student_name")
        ):

            dists = Distribution.objects.filter(
                student=st
            )

            b = (
                dists
                .filter(
                    benefit_type=BenefitType.BOOK
                )
                .aggregate(total=Sum("quantity"))
                ["total"]
                or 0
            )

            wb = (
                dists
                .filter(
                    benefit_type=BenefitType.WORKBOOK
                )
                .aggregate(total=Sum("quantity"))
                ["total"]
                or 0
            )

            sk = (
                dists
                .filter(
                    benefit_type=BenefitType.STUDY_KIT
                )
                .aggregate(total=Sum("quantity"))
                ["total"]
                or 0
            )

            lp = (
                dists
                .filter(
                    benefit_type=BenefitType.LAPTOP
                )
                .aggregate(total=Sum("quantity"))
                ["total"]
                or 0
            )

            writer.writerow(
                [
                    st.student_name,
                    st.admission_number,
                    st.current_class,
                    (
                        st.school.name
                        if st.school
                        else "General"
                    ),
                    b,
                    wb,
                    sk,
                    lp,
                    b + wb + sk + lp,
                ]
            )

    # --------------------------------------------------------
    # HISTORY
    # --------------------------------------------------------

    elif report_type == "history":

        response[
            "Content-Disposition"
        ] = (
            'attachment; '
            'filename="distribution_history.csv"'
        )

        writer = csv.writer(response)

        writer.writerow(
            [
                "Date",
                "Benefit Type",
                "Recipient Type",
                "Recipient / Student Name",
                "School",
                "Item Name",
                "Quantity",
                "Academic Year",
                "Issued By",
                "Remarks",
            ]
        )

        distributions = (
            Distribution.objects
            .select_related(
                "student",
                "school",
                "inventory_item",
                "study_kit",
            )
            .all()
            .order_by(
                "-distribution_date"
            )
        )

        for d in distributions:

            writer.writerow(
                [
                    d.distribution_date.strftime(
                        "%Y-%m-%d"
                    ),

                    d.get_benefit_type_display(),

                    d.get_recipient_type_display(),

                    (
                        d.student.student_name
                        if d.student
                        else (
                            d.school.name
                            if d.school
                            else "General"
                        )
                    ),

                    (
                        d.school.name
                        if d.school
                        else (
                            d.student.school.name
                            if (
                                d.student
                                and d.student.school
                            )
                            else ""
                        )
                    ),

                    (
                        d.inventory_item.item_name
                        if d.inventory_item
                        else (
                            d.study_kit.name
                            if d.study_kit
                            else "-"
                        )
                    ),

                    d.quantity,

                    d.academic_year,

                    d.issued_by,

                    d.remarks,
                ]
            )

    # --------------------------------------------------------
    # SCHOOL WISE
    # --------------------------------------------------------

    else:

        response[
            "Content-Disposition"
        ] = (
            'attachment; '
            'filename="school_distribution_report.csv"'
        )

        writer = csv.writer(response)

        writer.writerow(
            [
                "School Name",
                "UDISE Code",
                "District",
                "Books Distributed",
                "Workbooks",
                "Study Kits",
                "Laptops",
                "School Essentials",
                "Total Delivered",
            ]
        )

        for sc in (
            School.objects
            .all()
            .order_by("name")
        ):

            dists = Distribution.objects.filter(
                Q(school=sc)
                |
                Q(student__school=sc)
            )

            b = (
                dists
                .filter(
                    benefit_type=BenefitType.BOOK
                )
                .aggregate(total=Sum("quantity"))
                ["total"]
                or 0
            )

            wb = (
                dists
                .filter(
                    benefit_type=BenefitType.WORKBOOK
                )
                .aggregate(total=Sum("quantity"))
                ["total"]
                or 0
            )

            sk = (
                dists
                .filter(
                    benefit_type=BenefitType.STUDY_KIT
                )
                .aggregate(total=Sum("quantity"))
                ["total"]
                or 0
            )

            lp = (
                dists
                .filter(
                    benefit_type=BenefitType.LAPTOP
                )
                .aggregate(total=Sum("quantity"))
                ["total"]
                or 0
            )

            se = (
                dists
                .filter(
                    benefit_type=BenefitType.SCHOOL_ESSENTIAL
                )
                .aggregate(total=Sum("quantity"))
                ["total"]
                or 0
            )

            writer.writerow(
                [
                    sc.name,
                    sc.udise_code,
                    sc.district,
                    b,
                    wb,
                    sk,
                    lp,
                    se,
                    b + wb + sk + lp + se,
                ]
            )

    return response


# ============================================================================
# BULK UPLOAD: DISTRIBUTIONS (Feature 2)
# ============================================================================

def _normalize_dist_header(name):
    return str(name or "").strip().lower().replace(" ", "_").replace("-", "_")


def _is_school_roster_format(raw_rows):
    """
    Checks if the uploaded file is in School Profile / Roster format
    (contains multi-tier banners/headers for School, Taluk, UDISE, Class 4th-10th).
    """
    for r in raw_rows[:15]:
        row_str = " ".join(str(c or "").lower() for c in r)
        if ("school" in row_str or "hool" in row_str) and ("udise" in row_str or "taluk" in row_str):
            return True
        if "school profile" in row_str or "academic year" in row_str or "agastya" in row_str:
            return True
    return False


def _generate_temp_udise(school_name):
    import hashlib
    h = hashlib.md5(school_name.strip().lower().encode("utf-8")).hexdigest()[:10].upper()
    udise = f"SCH-{h}"
    counter = 1
    while School.objects.filter(udise_code=udise).exists():
        udise = f"SCH-{h[:6]}-{counter:03d}"
        counter += 1
    return udise


def _is_opn_notebook_format(raw_rows):
    """
    Checks if the uploaded sheet is an Operation Notebook (OPN) distribution sheet
    (e.g., contains School, Location, 80 Pg / 160 Pg, Total, Sent/Not Yet, Notebook Distributed).
    """
    for r in raw_rows[:10]:
        r_str = " ".join(str(c or "").lower() for c in r)
        if any(k in r_str for k in ("80 pg", "160 pg", "notebook distributed", "schoolwise notebook count", "total notebooks")):
            return True
        if ("school" in r_str or "name of the school" in r_str) and any(k in r_str for k in ("notebook", "sent/not yet", "photo posted", "pg", "completed", "sent")):
            return True
    return False


def _process_opn_notebook_distribution(raw_rows, request, selected_benefit_type=None, team_map=None):
    """
    Processes an Operation Notebook (OPN) distribution spreadsheet.
    Extracts schools, locations, notebook quantities, team issuers, and creates Distribution records.
    """
    if not team_map:
        team_map = {}

    benefit_type = selected_benefit_type or request.POST.get("benefit_type", "").strip().upper()
    if benefit_type not in dict(BenefitType.choices):
        benefit_type = BenefitType.BOOK

    academic_year = request.POST.get("academic_year", "").strip() or "2026-27"

    h_idx = 0
    for idx, r in enumerate(raw_rows[:6]):
        r_str = " ".join(str(c or "").lower() for c in r)
        if any(k in r_str for k in ("school", "location", "place", "admission", "benefit", "total")):
            h_idx = idx
            break

    headers = [str(c or "").strip().lower().replace(" ", "_").replace("-", "_").replace(".", "").replace("/", "_") for c in raw_rows[h_idx]]

    col_school = None
    col_loc = None
    col_total = None
    col_team = None
    col_status = None

    for i, h in enumerate(headers):
        if col_school is None and any(k in h for k in ("name_of_the_school", "school_name", "school", "institution")):
            col_school = i
        if col_loc is None and any(k in h for k in ("location", "place", "taluk", "district", "city")):
            col_loc = i
        if col_total is None and any(k in h for k in ("total", "quantity", "qty", "count", "strength", "notebooks", "total_notebooks")):
            col_total = i
        if col_team is None and any(k in h for k in ("team", "individual", "volunteer", "issued_by")):
            col_team = i
        if col_status is None and any(k in h for k in ("status", "distribution", "sent")):
            col_status = i

    footer_keywords = ("total", "requir", "stock", "short", "ordered", "delivered", "distributed", "difference", "gantt")
    created_distributions = []
    total_units = 0

    with transaction.atomic():
        for idx, r in enumerate(raw_rows[h_idx + 1:], start=h_idx + 2):
            if not r or not any(c is not None and str(c).strip() != "" for c in r):
                continue

            cell_str = " ".join(str(c or "").lower() for c in r[:4])
            if any(w in cell_str for w in footer_keywords):
                continue

            if any(str(c or "").strip() in ("1L", "2L", "4L", "SQ", "LB", "80 Pg", "160 Pg", "160 pg", "Total", "TOTAL") for c in r[:5]):
                continue

            school_name = str(r[col_school] or "").strip() if col_school is not None and col_school < len(r) else ""
            if not school_name or school_name.lower() in ("hubbali", "hubballi", "koppal", "belgaum", "belagavi", "dharwad", "bangalore"):
                continue

            loc = str(r[col_loc] or "").strip() if col_loc is not None and col_loc < len(r) else "Belagavi"
            if not loc:
                loc = "Belagavi"

            qty = 0
            if col_total is not None and col_total < len(r) and r[col_total] is not None:
                try:
                    qty = int(float(str(r[col_total]).replace(",", "").strip()))
                except Exception:
                    qty = 0

            if qty <= 0:
                nums = [int(float(c)) for c in r[3:9] if c is not None and str(c).strip().replace(".", "", 1).isdigit() and int(float(c)) > 0]
                qty = sum(nums) if nums else 1

            team = str(r[col_team] or "").strip() if col_team is not None and col_team < len(r) else ""
            status = str(r[col_status] or "").strip() if col_status is not None and col_status < len(r) else ""

            mapped_team, mapped_status = team_map.get(school_name.lower(), ("", ""))
            final_team = team or mapped_team or (request.user.username if request.user.is_authenticated else "Aequs Foundation")
            final_status = status or mapped_status or "Completed"

            school_obj = School.objects.filter(name__iexact=school_name).first()
            if not school_obj:
                school_obj = School.objects.filter(name__icontains=school_name).first()
            if not school_obj:
                school_obj = School.objects.create(
                    name=school_name,
                    udise_code=_generate_temp_udise(school_name),
                    district=loc,
                    taluk=loc,
                    status=School.Status.ACTIVE,
                )

            dist = Distribution.objects.create(
                benefit_type=benefit_type,
                recipient_type=RecipientType.SCHOOL,
                school=school_obj,
                academic_year=academic_year,
                quantity=qty,
                issued_by=final_team,
                location=loc,
                remarks=f"Operation Notebook (OPN {academic_year}): {qty:,} notebooks distributed. Status: {final_status}",
            )
            created_distributions.append(dist)
            total_units += qty

    if not created_distributions:
        raise ValueError("No school distribution records could be parsed from the uploaded spreadsheet.")

    return len(created_distributions), total_units, academic_year, benefit_type


def _process_school_profile_roster_distribution(raw_rows, request, selected_benefit_type=None):
    """
    Processes an Agastya / School Profile roster Excel file for distributions.
    Parses schools, UDISE codes, HM and Teacher contacts, grade strengths, and creates
    Distribution records for each school matching the enrolled student strength.
    """
    import re

    # 1. Detect Academic Year & Location from banners
    academic_year = request.POST.get("academic_year", "").strip() or "2025-26"
    default_location = "Dharwad"

    for r in raw_rows[:15]:
        row_str = " ".join(str(c or "").lower() for c in r)
        if "academic year" in row_str:
            m = re.search(r'20\d{2}[-\s/]\d{2,4}', row_str)
            if m:
                academic_year = m.group(0).replace(" ", "")
        if "location" in row_str:
            parts = [p.strip() for p in row_str.split(":") if p.strip()]
            if len(parts) > 1:
                default_location = parts[-1].title()

    # 2. Benefit Type selection
    benefit_type = selected_benefit_type or request.POST.get("benefit_type", "").strip().upper()
    if benefit_type not in dict(BenefitType.choices):
        benefit_type = BenefitType.STUDY_KIT

    # 3. Find Header Row
    header_idx = -1
    for idx, r in enumerate(raw_rows[:15]):
        row_str = " ".join(str(c or "").lower() for c in r)
        if ("school" in row_str and ("udise" in row_str or "taluk" in row_str)) or ("hool" in row_str and "taluk" in row_str):
            header_idx = idx
            break

    if header_idx == -1:
        for idx, r in enumerate(raw_rows[:10]):
            if sum(1 for c in r if c and isinstance(c, str) and len(c.strip()) > 1) >= 4:
                header_idx = idx
                break

    if header_idx == -1:
        header_idx = 0

    r0 = [str(c or "").strip() for c in raw_rows[header_idx]]
    col_map = {}
    mobiles = []
    for i, h in enumerate(r0):
        hl = h.lower()
        if "sl" in hl or ("no" in hl and "mobile" not in hl and "hool" not in hl):
            col_map.setdefault("sl", i)
        elif "school" in hl or "name" in hl and "teacher" not in hl and "hm" not in hl:
            col_map.setdefault("name", i)
        elif "taluk" in hl:
            col_map.setdefault("taluk", i)
        elif "udise" in hl or "dise" in hl:
            col_map.setdefault("udise", i)
        elif "hm" in hl or "headmaster" in hl:
            col_map.setdefault("hm", i)
        elif "teacher" in hl:
            col_map.setdefault("teacher", i)
        elif "div" in hl or "sec" in hl:
            col_map.setdefault("div", i)
        if "mobile" in hl or "phone" in hl or "contact" in hl:
            mobiles.append(i)

    if mobiles:
        col_map["hm_mob"] = mobiles[0]
        if len(mobiles) > 1:
            col_map["teacher_mob"] = mobiles[1]

    # Map grade start column
    start_grades_idx = 9
    for i, h in enumerate(r0):
        if "4" in h or "class" in h.lower():
            start_grades_idx = i
            break

    data_rows = raw_rows[header_idx + 1:]
    schools_data = []
    current_school = None
    seen_udises = {}

    def _to_int(val):
        if val is None or val == "" or val == "—" or val == "-":
            return 0
        try:
            return int(float(str(val).replace(",", "").strip()))
        except (ValueError, TypeError):
            return 0

    for r in data_rows:
        if not r or not any(c is not None and str(c).strip() != "" for c in r):
            continue

        first_cells_str = " ".join(str(c or "").lower() for c in r[:6])
        if any(w in first_cells_str for w in ("total", "grand total", "donor", "academic")):
            continue

        raw_name = str(r[col_map.get("name", 1)] or "").strip() if col_map.get("name") is not None and col_map.get("name") < len(r) else ""
        raw_udise = str(r[col_map.get("udise", 3)] or "").strip() if col_map.get("udise") is not None and col_map.get("udise") < len(r) else ""
        div_val = str(r[col_map.get("div", 8)] or "").strip() if col_map.get("div") is not None and col_map.get("div") < len(r) else ""

        def _get_grade_tuple(base_idx):
            if base_idx + 2 < len(r):
                b = _to_int(r[base_idx])
                g = _to_int(r[base_idx + 1])
                t = _to_int(r[base_idx + 2])
                if t == 0 and (b + g > 0):
                    t = b + g
                return [b, g, t]
            return [0, 0, 0]

        c4 = _get_grade_tuple(start_grades_idx)
        c5 = _get_grade_tuple(start_grades_idx + 3)
        c6 = _get_grade_tuple(start_grades_idx + 6)
        c7 = _get_grade_tuple(start_grades_idx + 9)
        c8 = _get_grade_tuple(start_grades_idx + 12)
        c9 = _get_grade_tuple(start_grades_idx + 15)
        c10 = _get_grade_tuple(start_grades_idx + 18)

        gt_idx = start_grades_idx + 21
        if gt_idx + 2 < len(r):
            gt_b = _to_int(r[gt_idx])
            gt_g = _to_int(r[gt_idx + 1])
            gt_t = _to_int(r[gt_idx + 2])
            if gt_t == 0:
                gt_t = sum(g[2] for g in (c4, c5, c6, c7, c8, c9, c10))
            gt = [gt_b, gt_g, gt_t]
        else:
            all_g = (c4, c5, c6, c7, c8, c9, c10)
            gt = [sum(g[0] for g in all_g), sum(g[1] for g in all_g), sum(g[2] for g in all_g)]

        div_info = {
            "division": div_val,
            "c4": c4, "c5": c5, "c6": c6, "c7": c7, "c8": c8, "c9": c9, "c10": c10,
            "gt": gt,
        }

        # Handle second division of the same school
        if (not raw_name or (current_school and raw_name == current_school.get("name"))) and current_school and div_val:
            current_school["divisions"].append(div_info)
            continue

        if not raw_name:
            continue

        if raw_udise in seen_udises:
            seen_udises[raw_udise] += 1
            raw_udise = f"{raw_udise}-{seen_udises[raw_udise]}"
        else:
            seen_udises[raw_udise] = 1

        taluk_val = str(r[col_map.get("taluk", 2)] or "").strip() if col_map.get("taluk") is not None and col_map.get("taluk") < len(r) else default_location
        hm_val = str(r[col_map.get("hm", 4)] or "").strip() if col_map.get("hm") is not None and col_map.get("hm") < len(r) else ""
        hm_mob_val = str(r[col_map.get("hm_mob", 5)] or "").strip() if col_map.get("hm_mob") is not None and col_map.get("hm_mob") < len(r) else ""
        teacher_val = str(r[col_map.get("teacher", 6)] or "").strip() if col_map.get("teacher") is not None and col_map.get("teacher") < len(r) else ""
        teacher_mob_val = str(r[col_map.get("teacher_mob", 7)] or "").strip() if col_map.get("teacher_mob") is not None and col_map.get("teacher_mob") < len(r) else ""

        current_school = {
            "name": raw_name,
            "taluk": taluk_val or default_location,
            "udise": raw_udise,
            "hm": hm_val,
            "hm_mob": hm_mob_val,
            "teacher": teacher_val,
            "teacher_mob": teacher_mob_val,
            "divisions": [div_info],
        }
        schools_data.append(current_school)

    if not schools_data:
        raise ValueError("No valid school records could be parsed from the uploaded spreadsheet.")

    created_distributions = []
    total_units = 0

    with transaction.atomic():
        for sc in schools_data:
            school_total = sum(d["gt"][2] for d in sc["divisions"])
            school_obj, _ = School.objects.update_or_create(
                udise_code=sc["udise"],
                defaults={
                    "name": sc["name"],
                    "taluk": sc["taluk"],
                    "district": "Dharwad",
                    "headmaster_name": sc["hm"],
                    "headmaster_phone": sc["hm_mob"],
                    "phone": sc["hm_mob"],
                    "student_strength": school_total,
                    "status": School.Status.ACTIVE,
                }
            )

            if sc["teacher"]:
                SchoolResource.objects.update_or_create(
                    school=school_obj,
                    resource_name=f"Science Teacher: {sc['teacher']}",
                    defaults={
                        "status": "Active",
                        "quantity": 1,
                        "details": f"Mobile: {sc['teacher_mob']}" if sc["teacher_mob"] else "Science Teacher",
                        "last_updated_note": "Imported from school profile distribution",
                    }
                )

            # Update grade strengths
            school_obj.grade_strengths.all().delete()
            grade_labels = [
                ("Class 4th", "c4", 4),
                ("Class 5th", "c5", 5),
                ("Class 6th (DLC)", "c6", 6),
                ("Class 7th", "c7", 7),
                ("Class 8th", "c8", 8),
                ("Class 9th", "c9", 9),
                ("Class 10th", "c10", 10),
            ]
            for g_label, g_key, g_ord in grade_labels:
                m = sum(d[g_key][0] for d in sc["divisions"])
                f = sum(d[g_key][1] for d in sc["divisions"])
                tot = sum(d[g_key][2] for d in sc["divisions"])
                if tot > 0:
                    GradeStrength.objects.create(
                        school=school_obj,
                        grade_level=g_label,
                        male_students=m,
                        female_students=f,
                        total_students=tot,
                        order=g_ord,
                    )

            # Create Distribution record for this school
            div_labels = [d["division"] for d in sc["divisions"] if d.get("division")]
            div_text = f" (Divisions: {', '.join(div_labels)})" if div_labels else ""
            dist_qty = school_total if school_total > 0 else 1

            issuer_name = sc["teacher"] or sc["hm"] or (request.user.username if request.user.is_authenticated else "Aequs Foundation")
            dist = Distribution.objects.create(
                benefit_type=benefit_type,
                recipient_type=RecipientType.SCHOOL,
                school=school_obj,
                academic_year=academic_year,
                quantity=dist_qty,
                issued_by=issuer_name,
                location=sc["taluk"] or default_location,
                remarks=f"School profile material distribution: {dist_qty} students enrolled{div_text}",
            )
            created_distributions.append(dist)
            total_units += dist_qty

    return len(created_distributions), total_units, academic_year, benefit_type


def bulk_upload_distributions(request):
    import io as _io
    if request.method != "POST":
        return render(request, "distributions/bulk_upload.html")

    uploaded_file = request.FILES.get("file") or request.FILES.get("csv_file")
    if not uploaded_file:
        messages.error(request, "Please select a CSV or Excel (.xlsx) file.")
        return redirect("distributions:bulk_upload")

    file_name = uploaded_file.name.lower()
    if not file_name.endswith((".csv", ".xlsx", ".xls")):
        messages.error(request, "Only CSV and Excel (.xlsx, .xls) files are supported.")
        return redirect("distributions:bulk_upload")

    try:
        team_map = {}
        if file_name.endswith((".xlsx", ".xls")):
            import openpyxl
            wb = openpyxl.load_workbook(uploaded_file, read_only=False, data_only=True)

            # Extract team/volunteer mapping from companion sheets if present (e.g. Notebook Distributed)
            for sname in wb.sheetnames:
                if any(k in sname.lower() for k in ("distributed", "team", "delivery")):
                    for r in wb[sname].iter_rows(values_only=True):
                        if r and len(r) > 4 and r[2]:
                            s_key = str(r[2]).strip().lower()
                            t_val = str(r[5] if len(r) > 5 and r[5] else "").strip()
                            st_val = str(r[4] if len(r) > 4 and r[4] else "").strip()
                            team_map[s_key] = (t_val, st_val)

            best_rows = []
            selected_sheet_rows = None
            for ws in wb.worksheets:
                cur = [list(r) for r in ws.iter_rows(values_only=True) if any(c is not None and str(c).strip() != "" for c in r)]
                if not cur:
                    continue
                if _is_school_roster_format(cur) or _is_opn_notebook_format(cur):
                    selected_sheet_rows = cur
                    break
                if len(cur) > len(best_rows):
                    best_rows = cur
            wb.close()
            raw_rows = selected_sheet_rows or best_rows
        else:
            decoded = ""
            content_bytes = uploaded_file.read()
            for enc in ("utf-8-sig", "utf-8", "cp1252", "iso-8859-1"):
                try:
                    decoded = content_bytes.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if not decoded:
                messages.error(request, "Could not decode CSV file. Please use UTF-8.")
                return redirect("distributions:bulk_upload")
            delim = "\t" if decoded[:2048].count("\t") > decoded[:2048].count(",") else ","
            reader = csv.reader(_io.StringIO(decoded), delimiter=delim)
            raw_rows = [list(r) for r in reader if any(c is not None and str(c).strip() != "" for c in r)]

        if not raw_rows:
            messages.error(request, "The uploaded file is empty.")
            return redirect("distributions:bulk_upload")

        # -------------------------------------------------------------
        # BRANCH 1: School Profile / Agastya Roster Format
        # -------------------------------------------------------------
        if _is_school_roster_format(raw_rows):
            created_count, total_qty, ay, b_type = _process_school_profile_roster_distribution(raw_rows, request)
            b_display = dict(BenefitType.choices).get(b_type, b_type)
            messages.success(
                request,
                f"Successfully parsed School Profile roster: created {created_count} school distribution records totaling {total_qty:,} units ({b_display}, Academic Year {ay})."
            )
            request.session["last_distribution_sheet"] = "agastya"
            return redirect(f"{reverse('distributions:list')}?view=tabular&sheet=agastya")

        # -------------------------------------------------------------
        # BRANCH 2: Operation Notebook (OPN) Distribution Format
        # -------------------------------------------------------------
        if _is_opn_notebook_format(raw_rows):
            created_count, total_qty, ay, b_type = _process_opn_notebook_distribution(raw_rows, request, team_map=team_map)
            b_display = dict(BenefitType.choices).get(b_type, b_type)
            messages.success(
                request,
                f"Successfully parsed Operation Notebook (OPN) distribution: created {created_count} school distribution records totaling {total_qty:,} units ({b_display}, Academic Year {ay})."
            )
            request.session["last_distribution_sheet"] = "opn"
            return redirect(f"{reverse('distributions:list')}?view=tabular&sheet=opn")

        # -------------------------------------------------------------
        # BRANCH 3: Flat / Tabular Columns Format
        # -------------------------------------------------------------
        header_idx = 0
        for idx, r in enumerate(raw_rows[:10]):
            row_str = " ".join(str(c or "").lower() for c in r)
            if any(k in row_str for k in ("benefit", "admission", "quantity", "school", "student", "total", "location", "place")):
                header_idx = idx
                break

        headers = [_normalize_dist_header(c) for c in raw_rows[header_idx]]
        rows = []
        for r in raw_rows[header_idx + 1:]:
            if not any(c is not None and str(c).strip() for c in r):
                continue
            rows.append({headers[i]: (str(r[i]).strip() if i < len(r) and r[i] is not None else "") for i in range(len(headers))})

        created, errors = 0, []
        default_benefit = request.POST.get("benefit_type", "").strip().upper() or BenefitType.BOOK
        with transaction.atomic():
            for idx, row in enumerate(rows, start=header_idx + 2):
                row_vals_str = " ".join(str(v or "").lower() for v in row.values())
                if any(w in row_vals_str for w in ("grand total", "total stock", "difference", "requirment", "ordered", "delivered")):
                    continue

                benefit_type = row.get("benefit_type") or row.get("benefit") or row.get("type") or default_benefit
                benefit_type = benefit_type.upper()
                if benefit_type not in dict(BenefitType.choices):
                    benefit_type = BenefitType.BOOK

                student = None
                admission = (
                    row.get("admission_number") or row.get("admission_no") or
                    row.get("adm_no") or row.get("roll_no") or row.get("student_id") or ""
                ).strip()
                if admission:
                    student = Student.objects.filter(admission_number=admission).first()
                if not student:
                    s_name = (row.get("student_name") or row.get("student") or "").strip()
                    if s_name:
                        student = Student.objects.filter(student_name__iexact=s_name).first()

                school = None
                school_name = (
                    row.get("school") or row.get("name_of_the_school") or
                    row.get("school_name") or row.get("institution") or
                    row.get("name_of_school") or ""
                ).strip()
                if not school_name:
                    for k, v in row.items():
                        if "school" in k and v and len(str(v)) > 2:
                            school_name = str(v).strip()
                            break

                loc = (
                    row.get("location") or row.get("place") or
                    row.get("taluk") or row.get("district") or ""
                ).strip()

                if school_name:
                    school = School.objects.filter(name__iexact=school_name).first()
                    if not school:
                        school = School.objects.filter(name__icontains=school_name).first()
                    if not school:
                        school = School.objects.create(
                            name=school_name,
                            udise_code=_generate_temp_udise(school_name),
                            district=loc or "Belagavi",
                            taluk=loc or "Belagavi",
                            status=School.Status.ACTIVE,
                        )

                if not student and not school:
                    if not any(len(str(v).strip()) > 0 for v in row.values()):
                        continue
                    errors.append(f"Row {idx}: No student or school specified")
                    continue

                qty_str = (
                    row.get("quantity") or row.get("total") or
                    row.get("qty") or row.get("count") or "1"
                ) or "1"
                try:
                    qty = int(float(str(qty_str).replace(",", "").strip()))
                except (ValueError, TypeError):
                    qty = 1

                Distribution.objects.create(
                    benefit_type=benefit_type,
                    recipient_type="STUDENT" if student else "SCHOOL",
                    student=student,
                    school=school or (student.school if student else None),
                    academic_year=row.get("academic_year", "2026-27") or "2026-27",
                    quantity=qty,
                    issued_by=row.get("issued_by") or row.get("team") or "",
                    location=loc or (school.taluk if school else ""),
                    remarks=row.get("remarks") or row.get("notes") or "",
                )
                created += 1

        msg = f"Bulk upload complete: {created} distributions created"
        if errors:
            msg += f", {len(errors)} errors"
        messages.success(request, msg)
        for e in errors[:10]:
            messages.warning(request, e)

        return redirect("distributions:list")

    except Exception as exc:
        messages.error(request, f"Upload failed: {exc}")
        return redirect("distributions:bulk_upload")

